#!/usr/bin/env python3
"""Rewrite the user draft into a method-focused AutoChecker technical report."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/测试用例驱动的静态代码检查器生成技术.docx"
OUTPUT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-技术报告版.docx"
ASSETS = ROOT / "docs/autochecker_method_report_assets"
PERF = ROOT / "AutoChecker_Performance.xlsx"

BLUE = "1F4E79"
MID_BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EDF4FA"
GRAY = "666666"
WHITE = "FFFFFF"


def east_asia_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_document(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for level, size, before, after in [(1, 16, 18, 10), (2, 13.5, 14, 7), (3, 11.5, 10, 5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "黑体"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else "365F91")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        ppr = style.element.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), str(level - 1))

    caption = doc.styles["Caption"]
    caption.font.name = "宋体"
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.color.rgb = RGBColor.from_string(GRAY)


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(text)
    east_asia_font(r, "黑体", 22, True, BLUE)


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    r = p.add_run(text)
    east_asia_font(r)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("• " + text)
    east_asia_font(r)
    return p


def caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_text(cell, value, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(str(value))
    east_asia_font(r, "宋体", size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def table(doc, headers, rows, widths=None, size=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    repeat_header(t.rows[0])
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], BLUE)
        cell_text(t.rows[0].cells[i], h, True, WHITE, WD_ALIGN_PARAGRAPH.CENTER, size)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, False, None,
                      WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT, size)
            if ridx % 2:
                shade(cells[i], PALE_BLUE)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def extract_overview():
    ASSETS.mkdir(exist_ok=True)
    target = ASSETS / "overview.jpeg"
    with zipfile.ZipFile(SOURCE) as archive:
        media = sorted(n for n in archive.namelist() if n.startswith("word/media/"))
        if not media:
            raise RuntimeError("源文档中没有找到 overview 图片")
        target.write_bytes(archive.read(media[0]))
    return target


def load_metrics():
    wb = load_workbook(PERF, data_only=True, read_only=True)
    ws = wb["Clang-Tidy"]
    details = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        details.append({
            "difficulty": row[0], "rule": row[1], "time": float(row[2]),
            "tokens": int(row[3]), "cost": float(row[4]),
            "pass": float(str(row[5]).strip("%")),
            "recall": float(str(row[6]).strip("%")),
        })
    ws = wb["Sheet2"]
    groups = []
    difficulty = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            difficulty = row[0]
        if row[1] == "Clang-Tidy":
            groups.append({
                "difficulty": difficulty, "time": float(row[3]), "tokens": int(row[4]),
                "cost": float(row[5]), "pass": float(str(row[6]).strip("%")),
                "recall": float(str(row[7]).strip("%")),
            })
    return details, groups


def load_knowledge_counts():
    with open(ROOT / "clang_tidy_collect/collect_clang_tidy_astMatchers/astMatchers.json", encoding="utf-8") as f:
        matcher_sections = json.load(f)
    matcher_count = sum(len(next(iter(section.values()))["matchers"]) for section in matcher_sections)
    with open(ROOT / "clang_tidy_collect/collect_clang_tidy_astMatchers_op/astMatchers_meta_op.json", encoding="utf-8") as f:
        matcher_ops = len(json.load(f))
    with open(ROOT / "clang_tidy_collect/collect_check_op/clang_tidy_check_op_dedup.json", encoding="utf-8") as f:
        check_ops = len(json.load(f))
    with open(ROOT / "clang_tidy_collect/collect_clang_tidy_ast_api/clang_tidy_ast_api.json", encoding="utf-8") as f:
        ast_api = json.load(f)
    ast_types = sum(len(items) for items in ast_api.values())
    ast_methods = sum(len(item.get("methods", [])) for items in ast_api.values() for item in items)
    return matcher_count, matcher_ops, check_ops, ast_types, ast_methods


def make_method_flow(path):
    fig, ax = plt.subplots(figsize=(11, 4.4))
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 4.5)
    ax.axis("off")
    data = [
        (.25, 1.55, 2.7, 1.3, "Stage 1\nPattern & Constraints", "#B4C7E7"),
        (4.15, 1.55, 2.7, 1.3, "Stage 2\nKnowledge Retrieval", "#9DC3E6"),
        (8.05, 1.55, 2.7, 1.3, "Stage 3\nSynthesis & Validation", "#A9D18E"),
    ]
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
    for x, y, w, h, text, color in data:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=.08",
                                    facecolor=color, edgecolor="#1F4E79", linewidth=1.6))
        ax.text(x+w/2, y+h/2, text, ha="center", va="center", fontsize=12, weight="bold")
    ax.add_patch(FancyArrowPatch((2.95, 2.2), (4.15, 2.2), arrowstyle="-|>", mutation_scale=16,
                                 color="#1F4E79", linewidth=1.5))
    ax.add_patch(FancyArrowPatch((6.85, 2.2), (8.05, 2.2), arrowstyle="-|>", mutation_scale=16,
                                 color="#1F4E79", linewidth=1.5))
    ax.add_patch(FancyArrowPatch((9.4, 1.5), (1.6, 1.05), connectionstyle="arc3,rad=.18",
                                 arrowstyle="-|>", mutation_scale=15, color="#7030A0", linewidth=1.5))
    ax.text(5.5, .35, "Failed test cases feed the next refinement round", ha="center",
            fontsize=10.5, color="#7030A0")
    ax.text(1.6, 3.45, "rule + test cases + AST", ha="center", color="#666666", fontsize=10)
    ax.text(5.5, 3.45, "API signatures + Meta-OP snippets", ha="center", color="#666666", fontsize=10)
    ax.text(9.4, 3.45, "compilable checker + expected behavior", ha="center", color="#666666", fontsize=10)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_effect_chart(path, groups):
    labels = ["Easy", "Medium", "Hard"]
    pass_rates = [x["pass"] for x in groups]
    recalls = [x["recall"] for x in groups]
    x = list(range(3))
    width = .34
    fig, ax = plt.subplots(figsize=(8.8, 4.4))
    bars1 = ax.bar([v-width/2 for v in x], pass_rates, width, color="#5B9BD5", label="Test pass rate")
    bars2 = ax.bar([v+width/2 for v in x], recalls, width, color="#70AD47", label="Violation recall")
    ax.set_xticks(x, labels)
    ax.set_ylim(94, 101.5)
    ax.set_ylabel("Percent (%)")
    ax.set_title("Checker generation results by rule difficulty", color="#1F4E79", weight="bold")
    ax.grid(axis="y", linestyle="--", alpha=.25)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(loc="lower left")
    for bars in (bars1, bars2):
        for b in bars:
            ax.text(b.get_x()+b.get_width()/2, b.get_height()+.18, f"{b.get_height():.1f}",
                    ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_picture(doc, image_path, width, label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    caption(doc, label)


def build():
    overview = extract_overview()
    details, groups = load_metrics()
    knowledge = load_knowledge_counts()
    make_method_flow(ASSETS / "three_stage_flow.png")
    make_effect_chart(ASSETS / "technical_effect.png", groups)

    doc = Document()
    setup_document(doc)
    title(doc, "测试用例驱动的静态代码检查器生成技术")

    doc.add_heading("1 背景技术", level=1)
    doc.add_heading("1.1 静态代码检查器的定制需求", level=2)
    body(doc, "静态代码检查是在不运行程序的情况下分析源代码，并据此发现编码规范违例、潜在缺陷和安全风险。它能够在编码和集成阶段提前暴露问题，因而被广泛用于代码评审、持续集成和软件质量控制。通用静态分析工具通常提供一批内置检查规则，但这些规则主要面向共性问题，难以完整覆盖企业内部规范、行业编码标准以及具体项目的设计约束。")
    body(doc, "在实际项目中，同一条规则往往还包含适用范围、例外条件和特定的诊断要求。例如，某种语法结构在一般上下文中应当报告，但在宏展开、模板实例化或特定类型条件下需要排除。此类要求很难直接由通用规则表达，因此需要针对项目开发专用检查器。")

    doc.add_heading("1.2 现有方法存在的问题", level=2)
    body(doc, "人工开发检查器需要先把自然语言规则转换为程序结构上的判定条件，再选择合适的抽象语法树节点和分析接口实现这些条件。开发人员不仅要熟悉目标语言，还要理解编译器前端、节点层次、匹配接口、源代码位置和诊断机制。规则稍有变化，就可能需要重新调整节点匹配和上下文判断。")
    body(doc, "直接使用大语言模型生成检查器虽然可以降低编码工作量，但仍有三个问题。第一，模型可能使用不存在或签名不匹配的接口，生成结果无法编译；第二，规则文本通常不会列出全部语法变体，模型容易遗漏需要报告的情况；第三，模型可能把示例中的偶然特征当作规则条件，导致对合规代码产生误报。因此，仅给出规则描述并一次性生成代码，无法稳定得到可用的检查器。")

    doc.add_heading("1.3 测试用例驱动生成的必要性", level=2)
    body(doc, "规则描述说明检查目标，测试用例则给出可以直接验证的行为边界。违规用例表示检查器应当报告的代码模式，合规用例表示检查器应当保持静默的情形。两类用例共同限定了检查器的检测范围：前者用于补充召回范围，后者用于形成排除条件。")
    body(doc, "测试用例的作用不应局限在最终验收。AutoChecker在生成开始时利用违规用例识别代码模式和AST结构，在知识检索时利用样例节点筛选接口，在候选代码产生后利用完整测试套件发现漏报和误报，并把失败用例重新送入生成过程。由此，检查器的形成过程由可执行样例持续约束，而不是完全依赖模型对自然语言的自由解释。")

    doc.add_heading("2 测试用例驱动的静态代码检查器生成总体框架", level=1)
    doc.add_heading("2.1 方法输入与输出", level=2)
    body(doc, "方法输入由规则描述和测试套件组成。规则描述给出检测对象、违规条件、允许情形和诊断意图；测试套件由违规用例和合规用例构成，每个用例都带有明确的预期结果。对于当前C/C++实现，测试用例以源文件形式组织，预期产生诊断的代码位置由测试标记给出。")
    body(doc, "方法输出是能够接入目标静态分析框架的检查器源代码。输出结果需要同时满足两个条件：一是通过目标工程的编译，接口调用、类型和依赖关系正确；二是在测试套件上的运行结果与预期一致，即违规用例得到报告，合规用例不产生报告。当前系统以clang-tidy作为C/C++检查器的承载框架。")
    table(doc, ["要素", "内容", "在生成过程中的作用"], [
        ("规则描述", "规则目标、适用范围、违规条件和允许情形", "提供检测语义和判断边界"),
        ("违规用例", "预期产生检查报告的代码", "提供目标代码模式并检验漏报"),
        ("合规用例", "预期不产生检查报告的代码", "提供排除条件并检验误报"),
        ("输出结果", "可编译、可运行的检查器源代码", "作为编译验证和测试验证对象"),
    ], widths=[2.6, 6.0, 6.3])

    doc.add_heading("2.2 方法总体流程", level=2)
    body(doc, "检查器生成包括三个连续阶段。阶段1分析规则和测试代码，提取AST特征，生成检测逻辑集合及检查器约束；阶段2以检测子逻辑为查询，从API数据库和元操作数据库中检索相关接口与代码片段；阶段3综合规则、测试用例、AST、检测逻辑和检索结果生成候选检查器，并通过编译和测试反馈持续修正。")
    add_picture(doc, overview, 6.0, "图 1  测试用例驱动的静态代码检查器生成总体框架")
    body(doc, "图1中的验证环节服务于前三个生成阶段。编译失败时，反馈内容主要用于修正候选代码和补充接口知识；测试失败时，反馈内容用于调整检测逻辑和检查器约束。因此，从检查器生成的角度看，整体过程仍由三个阶段构成，验证是第三阶段中的反馈机制。")
    add_picture(doc, ASSETS / "three_stage_flow.png", 5.9, "图 2  三阶段生成过程及测试反馈关系")

    doc.add_heading("2.3 测试用例的驱动作用", level=2)
    body(doc, "测试用例贯穿三个阶段，但在不同阶段承担的作用不同。阶段1关注样例所体现的程序结构，通过违规与合规模式确定目标节点和排除条件；阶段2关注样例实际包含的AST节点类型，用这些结构信息过滤仅在语义上相似、但无法作用于当前代码的API；阶段3关注检查器的外部行为，根据每个用例的期望结果判断是否存在漏报或误报。")
    table(doc, ["阶段", "测试用例提供的信息", "形成的约束"], [
        ("阶段1", "代码模式、节点层次、违规与合规差异", "目标节点、结构条件和排除条件"),
        ("阶段2", "样例中实际出现的AST节点类型", "检索结果必须与样例结构相关"),
        ("阶段3", "预期报告或预期静默", "候选检查器必须满足完整测试套件"),
    ], widths=[2.2, 6.4, 6.3])
    body(doc, "这种组织方式把规则文本和测试用例看作同一份规格的两种表达。规则文本负责概括，用例负责落地；当二者存在信息缺口时，测试结果会直接暴露缺口，并触发下一轮逻辑调整。")

    doc.add_heading("3 测试用例驱动的静态代码检查器生成方法", level=1)
    doc.add_heading("3.1 基于代码模式生成检测逻辑集合和检查器约束", level=2)
    doc.add_heading("3.1.1 测试用例代码模式分析", level=3)
    body(doc, "阶段1首先从测试套件中选择一个违规用例作为生成种子。选择违规用例的原因是初始检查器至少需要具备识别一个确定违例的能力，才能进入后续的全量测试和增量修正。若当前用例在限定轮次内不能形成有效检查器，则选择其他违规用例重新开始。")
    body(doc, "代码模式分析同时参考规则描述和种子用例。规则描述用于判断哪些结构具有规则意义，测试代码用于确定这些结构在源程序中的具体形式。分析内容包括目标程序实体、目标实体与上下文节点的关系、必须满足的属性条件、需要排除的合法形式以及适合输出诊断的位置。分析时不直接使用示例中的变量名和字面量作为规则条件，除非规则本身明确要求。")
    body(doc, "初始阶段以单个违规用例建立可以运行的检查器，完整的规则边界由后续测试反馈逐步补充。这样可以先解决“能否识别一个明确违例”的问题，再处理多种语法形式和例外条件，避免在首次生成时同时引入过多相互影响的约束。")

    doc.add_heading("3.1.2 抽象语法树特征提取", level=3)
    body(doc, "源代码表面形式不能完整反映程序结构。例如，条件表达式可能包含隐式类型转换，函数调用可能通过不同语法形式引用同一声明。为获得与静态分析框架一致的结构信息，系统对测试用例进行前端解析，得到抽象语法树。")
    body(doc, "提取过程保留与当前测试代码相关的AST分支，并从中识别节点类型。节点类型集合用于说明样例中实际存在的声明、语句和表达式类型；AST文本则保留父子关系、源代码位置和类型信息。两者分别用于检测逻辑生成和后续知识筛选。")
    body(doc, "以“条件表达式中不得直接使用赋值”为例，源码中的赋值符号会对应赋值运算节点，该节点又位于条件语句的条件子树中。检测逻辑不仅要识别赋值运算，还要描述它与条件语句之间的结构关系。AST为这种关系提供了比源码字符串更稳定的表达。")

    doc.add_heading("3.1.3 检测逻辑集合生成", level=3)
    body(doc, "系统将完整的规则判定拆成若干可以单独实现的检测步骤。检测逻辑分为模式定位逻辑和检查验证逻辑。模式定位逻辑负责在AST中筛选候选节点并建立节点绑定；检查验证逻辑负责读取绑定节点的属性、检查上下文、执行排除判断并在合适位置输出诊断。")
    table(doc, ["逻辑类型", "主要内容", "典型逻辑"], [
        ("模式定位逻辑", "目标节点、父子关系、祖先关系和节点绑定", "定位条件表达式中的赋值运算节点"),
        ("检查验证逻辑", "属性读取、语义判断、排除条件和诊断位置", "确认节点位置有效并排除允许的比较表达式"),
    ], widths=[3.0, 6.0, 5.9])
    body(doc, "逻辑生成结果采用结构化形式保存，两类逻辑分别形成有序列表。每个步骤只描述一个明确操作，使其既可以作为代码生成依据，也可以作为知识库检索查询。生成结果在进入下一阶段前需要通过格式和字段检查；结构不完整时重新生成，避免错误逻辑继续传递。")

    doc.add_heading("3.1.4 检查器约束表示", level=3)
    body(doc, "检测逻辑集合可以进一步概括为检查器约束表示，即overview图中的Checker Schema。它不是对测试代码的复制，而是对检查器必须实现的结构条件进行整理。主要内容包括目标节点、需要绑定的节点、目标节点之间的结构关系、上下文条件、排除条件和诊断位置。")
    body(doc, "当前实现中，检查器约束没有作为单独文件输出，而是分布在模式定位逻辑、检查验证逻辑、样例AST和生成提示中。Checker Schema用于描述这些约束的逻辑结构。这样的表述既保留了方法层面的统一模型，也与当前代码中两类检测逻辑的实际表示一致。")
    body(doc, "检查器约束的作用是缩小生成空间。模型不再仅根据一句规则描述猜测完整实现，而是针对已经明确的节点、关系和条件生成代码。后续测试中发现的新语法形式或排除条件，也可以先更新约束，再据此重新检索和生成。")

    doc.add_heading("3.2 基于检测逻辑检索API上下文", level=2)
    doc.add_heading("3.2.1 静态分析知识的组织", level=3)
    body(doc, "检查器代码需要调用目标静态分析框架提供的节点匹配、属性访问、上下文遍历和诊断接口。仅依赖模型记忆容易出现接口名称、参数类型或组合方式错误。阶段2把框架相关知识预先整理成可检索的API知识和元操作知识，在生成前为每条检测逻辑补充可用的编程上下文。")
    body(doc, "API知识描述单个接口能够完成的操作，并保留方法签名、参数和返回类型。元操作知识描述一项较完整的检查动作，并给出来自既有检查器的实现片段。前者解决“接口是什么”，后者解决“接口通常怎样组合”。根据检测过程的不同位置，知识被组织为模式匹配API、AST节点操作API、匹配元操作和检查验证元操作四类。")
    m, mo, co, at, am = knowledge
    table(doc, ["知识类型", "规模", "内容", "用途"], [
        ("模式匹配API", f"{m}个", "节点匹配、条件收窄和AST遍历接口", "实现模式定位逻辑"),
        ("匹配元操作", f"{mo}条", "匹配语义及对应组合片段", "提供可复用的匹配方式"),
        ("检查验证元操作", f"{co}条", "节点获取、条件判断和诊断片段", "实现检查验证逻辑"),
        ("AST节点操作API", f"{at}个类型、{am}个方法", "AST类、结构体及其方法签名", "读取节点属性和上下文"),
    ], widths=[3.1, 3.0, 5.5, 3.4])
    body(doc, "表中规模由当前仓库的知识库文件统计得到。较大的知识集合提高了可覆盖的接口范围，但也使直接把全部知识放入生成上下文变得不可行，因此需要按照检测子逻辑进行检索。")

    doc.add_heading("3.2.2 基于检测子逻辑的语义检索", level=3)
    body(doc, "阶段1产生的每条检测逻辑分别作为检索查询。模式定位逻辑在模式匹配API和匹配元操作中检索，检查验证逻辑在AST节点操作API和检查验证元操作中检索。与使用完整规则描述进行一次查询相比，细粒度逻辑能够把不同目的的接口分开检索，减少高相似但用途不一致的结果。")
    body(doc, "检索时，检测逻辑和知识描述被转换为归一化语义向量，向量内积作为余弦相似度。每条逻辑取得分最高的一组候选，多条逻辑的结果合并并去重。对于元操作，检索对象是框架无关的操作描述，命中后再返回与目标框架对应的实现片段。")
    body(doc, "这种做法把规则语义与框架代码连接起来。检测逻辑仍然使用“获取被调用函数”“判断节点是否位于条件表达式”等功能性描述，而检索结果提供具体接口和代码形式。生成阶段因此不需要从自然语言直接推断全部框架细节。")

    doc.add_heading("3.2.3 基于AST特征的检索结果筛选", level=3)
    body(doc, "语义相似度只能反映文本含义是否接近，不能保证检索结果适用于当前样例。例如，多个接口都可能描述“获取声明名称”，但它们面向的AST节点类型不同。系统使用阶段1提取的节点类型集合对候选结果进行第二次筛选。")
    body(doc, "候选知识中出现的节点类型与样例AST一致时提高排序；C++特有节点与其通用节点形式一致时保留较低的相关分；父子遍历、祖先查找、逻辑组合等通用匹配操作保留基础分。若候选知识没有出现可识别的节点类型，则保持原有语义排序，避免因样例AST信息不完整而错误删除结果。")
    body(doc, "经过语义检索和AST筛选后，每条检测逻辑获得少量、与当前程序结构更一致的接口说明和实现片段。这些结果共同构成阶段3使用的API上下文。")

    doc.add_heading("3.2.4 API上下文的统一组织", level=3)
    body(doc, "检索结果按照模式定位和检查验证两个部分组织。模式定位部分包含候选节点匹配器、结构关系匹配器和绑定方式；检查验证部分包含节点获取、类型与属性读取、上下文访问、源码位置处理和诊断输出方式。接口说明与元操作片段同时保留，模型既能看到合法签名，也能看到接口之间的常见组合。")
    body(doc, "API上下文只提供完成当前检测逻辑所需的信息，不包含与样例结构无关的大量框架文档。这样既控制了输入长度，也降低了相似接口之间的干扰。若后续编译错误表明某项知识仍然不足，可以根据错误分析结果进行补充检索。")

    doc.add_heading("3.3 基于检查器约束的候选检查器合成与验证", level=2)
    doc.add_heading("3.3.1 多源生成上下文构建", level=3)
    body(doc, "阶段3将规则描述、测试用例、样例AST、检测逻辑、检查器约束、API上下文和目标框架模板组织为一次候选生成任务。各部分解决的问题不同：规则描述限定语义目标，测试用例提供具体边界，AST说明实际程序结构，检测逻辑给出处理顺序，API上下文约束接口使用，模板确定检查器需要遵循的工程结构。")
    body(doc, "初始生成以一个违规种子为主要测试依据。进入增强阶段后，生成上下文还会加入当前检查器、已通过用例以及当前失败用例。已通过用例用于保护已经满足的行为，失败用例用于指出当前实现尚未覆盖或排除的代码形式。")

    doc.add_heading("3.3.2 候选检查器合成", level=3)
    body(doc, "生成模型在目标框架提供的基础模板上形成完整候选检查器。候选内容至少包括模式注册和命中检查两部分：模式注册把阶段1的目标节点与结构约束转换为节点匹配表达式，并为检查阶段需要的节点建立绑定；命中检查读取绑定节点，完成属性、上下文和排除判断，最后在指定源码位置输出诊断。")
    body(doc, "检索到的元操作片段可以直接提供局部实现参考，但最终结果不是简单拼接片段。生成时还需要处理绑定名称一致性、接口参数类型、命名空间、头文件依赖和诊断位置等关系。当前C/C++实现一次生成完整的实现文件和头文件，只有两部分都能够被正确解析时才进入编译验证。")

    doc.add_heading("3.3.3 编译反馈修正", level=3)
    body(doc, "候选检查器首先接入目标静态分析框架进行编译。编译验证能够发现自然语言和测试运行无法直接暴露的问题，包括接口不存在、参数类型不匹配、命名空间错误、缺少头文件、绑定节点类型不一致以及方法签名不符合框架要求。")
    body(doc, "原始构建日志包含大量进度信息和编译命令，直接用于错误分析会掩盖真正的诊断。系统先清理与错误无关的内容，再将候选代码和有效诊断交给错误分析过程。错误分析给出具体修正步骤，并指出还需要检索的接口或代码形式。补充知识与修正步骤一起用于生成新版本，然后重新编译。")
    body(doc, "编译反馈解决的是候选检查器“能否作为程序工作”的问题。只有通过编译的候选结果才进入测试验证，避免把接口和语法错误与规则行为错误混在同一轮修正中。")

    doc.add_heading("3.3.4 测试反馈增强", level=3)
    body(doc, "候选检查器通过编译后，先在当前种子用例上运行。能够在种子违例位置产生报告，说明初始检测路径有效；随后运行完整测试套件，根据每个用例的预期行为划分通过用例和失败用例。")
    body(doc, "违规用例没有产生报告时，失败类型为漏报。这通常说明目标节点覆盖不足、上下文路径不完整或检测条件过严。系统根据漏报用例重新生成检测逻辑，补充新的语法形式和节点关系，并检索相应接口。合规用例产生报告时，失败类型为误报。这通常说明匹配范围过宽或缺少排除条件，增强过程会增加类型、作用域、父子关系或语义条件限制。")
    body(doc, "两类失败使用不同的反馈方向，但都同时参考当前检查器和已经通过的用例。修正某个失败用例后，系统重新运行全部测试用例，而不是只验证刚刚修正的样例。全量回归能够发现新增匹配条件造成的误报，也能发现新增排除条件导致的漏报。")

    doc.add_heading("3.3.5 生成闭环与终止条件", level=3)
    body(doc, "编译反馈和测试反馈构成两个相互衔接的闭环。编译失败时，候选代码在保持检测目标不变的情况下进行接口级修正；测试失败时，检测逻辑、检查器约束和API上下文根据失败用例重新组织，再产生新的候选代码。")
    body(doc, "当完整测试套件中的违规用例均得到报告、合规用例均保持静默时，当前检查器作为最终结果输出。为避免个别样例造成无界循环，每个种子、编译修复和用例增强过程都受到轮次限制；达到限制后可以更换种子或记录未解决用例。这里的轮次限制用于控制生成过程，不改变最终检查器应满足的行为要求。")

    doc.add_heading("4 技术效果", level=1)
    doc.add_heading("4.1 检查器生成效果", level=2)
    mean_time = sum(x["time"] for x in details) / len(details)
    mean_tokens = sum(x["tokens"] for x in details) / len(details)
    mean_cost = sum(x["cost"] for x in details) / len(details)
    mean_pass = sum(x["pass"] for x in details) / len(details)
    mean_recall = sum(x["recall"] for x in details) / len(details)
    body(doc, "现有性能记录中包含10条C/C++规则，每条规则配有20个测试用例，其中10个为违规用例、10个为合规用例，共计200个测试用例。统计结果见下表。该数据用于说明当前方法在已有测试套件上的生成效果。")
    table(doc, ["规则数量", "测试用例数量", "平均测试通过率", "违规用例召回率", "平均生成耗时", "平均Token消耗", "平均成本"], [[
        len(details), len(details) * 20, f"{mean_pass:.1f}%", f"{mean_recall:.1f}%",
        f"{mean_time:.2f}秒", f"{mean_tokens:,.0f}", f"{mean_cost:.4f}元"
    ]], widths=[1.8, 2.2, 2.4, 2.3, 2.3, 2.4, 2.0], size=8.2)
    body(doc, "10条规则中有9条通过全部测试用例。按样例数量计算，200个测试用例中有199个结果符合预期；100个违规用例中有99个被正确报告。未达到满分的规则为no-same-name-as-global-variable，其测试通过率为95%，违规用例召回率为90%，表现为一个违规代码变体未被覆盖。")
    body(doc, "这一结果说明，编译验证能够消除大部分接口和结构错误，完整测试套件的持续反馈能够进一步修正检测范围。对于已有规则集合，最终错误主要表现为少量语义变体漏报，而不是大范围编译失败或合规代码误报。")

    doc.add_heading("4.2 不同难度规则的生成结果", level=2)
    table(doc, ["规则难度", "平均测试通过率", "违规用例召回率", "平均耗时", "平均Token消耗"], [
        (x["difficulty"], f"{x['pass']:.1f}%", f"{x['recall']:.1f}%", f"{x['time']:.2f}秒", f"{x['tokens']:,}")
        for x in groups
    ], widths=[2.6, 3.0, 3.0, 3.0, 3.4])
    add_picture(doc, ASSETS / "technical_effect.png", 5.9, "图 3  不同难度规则的测试通过率与违规用例召回率")
    body(doc, "简单规则的平均测试通过率和违规用例召回率均为100%。中等规则的两项指标分别为98.8%和97.5%，差异来自同名变量规则中的一个漏报样例。困难规则在当前测试集合中同样达到100%，但平均耗时和Token消耗明显增加：平均耗时由简单规则的994.02秒增加到2226.11秒，平均Token消耗由111,937增加到271,596。")
    body(doc, "因此，在当前数据范围内，规则复杂度主要影响生成和修正成本。较复杂规则需要处理更多上下文关系，通常会经历更长的候选生成和反馈过程；测试用例驱动的闭环仍能够把最终结果约束到较高的测试通过水平。")

    doc.core_properties.title = "测试用例驱动的静态代码检查器生成技术"
    doc.core_properties.subject = "AutoChecker静态代码检查器生成方法技术报告"
    doc.core_properties.author = "AutoChecker项目组"
    doc.save(OUTPUT)
    return {
        "source": str(SOURCE),
        "output": str(OUTPUT),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "figures": len(doc.inline_shapes),
        "rules": len(details),
        "tests": len(details) * 20,
        "pass_rate": mean_pass,
        "recall": mean_recall,
    }


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False, indent=2))
