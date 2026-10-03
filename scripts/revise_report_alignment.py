#!/usr/bin/env python3
"""Align report terminology and Word heading semantics in place."""

import os
import tempfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一二章修改版.docx"

REPLACEMENTS = {
    "如图1所示，检查器生成过程由三个阶段组成。第一阶段以违规用例为种子，结合规则描述和测试代码的抽象语法树提取待检查的代码模式，并将检查过程拆分为模式定位逻辑和检查验证逻辑，在此基础上形成目标节点、节点关系、上下文条件、排除条件和诊断位置等检查器结构约束。第二阶段以各项检测子逻辑为查询，从模式匹配API、AST节点操作API、匹配元操作和检查验证元操作等知识库中检索相关接口及实现片段，并结合测试用例中实际出现的AST节点类型对检索结果进行筛选，形成与当前代码模式相匹配的API上下文。第三阶段综合规则描述、测试用例、抽象语法树、检测逻辑、结构约束和检索结果，在目标检查器模板基础上生成完整的候选检查器。":
    "如图1所示，检查器生成过程由三个阶段组成。第一阶段以具有代表性的违规用例为种子，结合规则描述和样例抽象语法树（AST）提取待检查的代码模式，并将检查过程拆分为模式定位逻辑和检查验证逻辑，在此基础上形成检测逻辑集合和检查器结构约束。第二阶段以检测逻辑为基本查询单元，分别从API数据库和Meta OP数据库中召回相关知识。其中，模式定位逻辑查询模式匹配API子集和匹配Meta OP子集，检查验证逻辑查询AST节点操作API子集和检查验证Meta OP子集；在语义召回的基础上，结合样例AST进行结构相关性重排，形成由接口说明、方法签名和可复用实现片段组成的知识上下文。第三阶段综合规则描述、检测逻辑集合、检查器结构约束、知识上下文和样例AST，生成完整的候选检查器。",

    "候选检查器生成后依次进行编译验证和测试验证。编译失败时，根据错误信息修正接口调用和代码结构，并补充检索所需的API知识；编译通过后，使用完整测试套件检查候选结果。违规用例未产生告警时，利用该用例补充检测模式；合规用例产生告警时，利用该用例增加排除条件。每次修正后重新执行测试，使新增逻辑不破坏已经通过的用例。":
    "候选检查器生成后依次进行编译验证和测试验证。编译失败时，根据错误信息修正接口调用和代码结构；当现有知识不足时，从API数据库或Meta OP数据库中补充检索相关知识。编译通过后，使用完整测试套件检查候选结果：违规用例未产生告警时判定为漏报，并据此补充检测模式；合规用例产生告警时判定为误报，并据此增加排除条件。测试失败产生的修正要求反馈至第一阶段，更新后的检测逻辑依次重新进入知识检索、候选生成和验证过程。每次修正后均重新执行完整测试套件，避免新增逻辑破坏已经通过的用例。",

    "截至中期检查阶段，项目已完成上述三阶段生成流程以及由编译反馈和测试反馈构成的迭代闭环，完成了模式匹配API、AST节点操作API、匹配元操作和检查验证元操作等知识的收集与组织。现有阶段性数据包含10条C/C++规则和200个测试用例，生成检查器的平均测试通过率为99.5%，违规用例召回率为99.0%。后续将继续扩大规则验证范围，并重点完善复杂规则的生成稳定性和处理效率。":
    "截至中期检查阶段，项目已完成上述三阶段生成流程以及由编译反馈和测试反馈构成的迭代闭环，并完成API数据库和Meta OP数据库的初步构建。当前纳入阶段性验证的对象包括10条C/C++规则和200个测试用例，生成检查器的平均测试通过率为99.5%，违规用例召回率为99.0%。上述结果用于反映现阶段生成检查器在既定测试套件上的规则符合性；后续将继续扩大规则验证范围，并通过独立测试样例进一步评估复杂规则的生成稳定性和泛化能力。",
}


def set_heading_styles(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in {"1 技术概述", "2 背景"} or text.startswith("3 测试用例驱动"):
            paragraph.style = "Heading 1"
        elif text.startswith(("3.1 ", "3.2 ", "3.3 ", "3.4 ")):
            paragraph.style = "Heading 2"
        elif text.startswith(("3.1.", "3.2.", "3.3.", "3.4.")):
            paragraph.style = "Heading 3"


def main():
    doc = Document(DOCUMENT)
    remaining = dict(REPLACEMENTS)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in remaining:
            paragraph.text = remaining.pop(text)
    if remaining:
        raise RuntimeError(f"未找到待替换段落：{list(remaining)}")

    set_heading_styles(doc)

    fd, temp_name = tempfile.mkstemp(
        prefix=".autochecker-alignment-", suffix=".docx", dir=str(DOCUMENT.parent)
    )
    os.close(fd)
    try:
        doc.save(temp_name)
        os.replace(temp_name, DOCUMENT)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    print(DOCUMENT)


if __name__ == "__main__":
    main()
