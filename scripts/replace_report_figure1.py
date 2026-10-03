#!/usr/bin/env python3
"""Replace Figure 1 in the AutoChecker midterm-report DOCX in place."""

import os
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一二章修改版.docx"
FIGURE = ROOT / "docs/专利图v6-技术报告版-v2.png"


def main():
    if not DOCUMENT.exists():
        raise FileNotFoundError(DOCUMENT)
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    doc = Document(DOCUMENT)
    caption_index = next(
        i for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip().startswith("图1")
    )
    figure_paragraph = next(
        paragraph for paragraph in reversed(doc.paragraphs[:caption_index])
        if paragraph._p.xpath(".//w:drawing")
    )

    figure_paragraph.clear()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = figure_paragraph.add_run()
    run.add_picture(str(FIGURE), width=Cm(15.37))

    fd, temp_name = tempfile.mkstemp(
        prefix=".autochecker-figure1-", suffix=".docx", dir=str(DOCUMENT.parent)
    )
    os.close(fd)
    try:
        doc.save(temp_name)
        os.replace(temp_name, DOCUMENT)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)

    print(DOCUMENT)


if __name__ == "__main__":
    main()
