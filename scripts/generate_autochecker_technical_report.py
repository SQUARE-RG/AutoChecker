#!/usr/bin/env python3
"""Generate the AutoChecker clang-tidy checker-generation technical report."""

from __future__ import annotations

import json
import math
import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from openpyxl import load_workbook
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "autochecker_report_assets"
OUT = DOCS / "测试用例驱动的静态代码检查器生成技术.docx"
PERF = ROOT / "AutoChecker_Performance.xlsx"

BLUE = "1F4E79"
BLUE2 = "2F75B5"
LIGHT_BLUE = "D9EAF7"
VERY_LIGHT_BLUE = "EEF5FA"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
RED = "C00000"
GREEN = "548235"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_col_widths(table, widths_cm) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_shading(hdr.cells[i], BLUE)
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=font_size,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        hdr.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr.cells[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size,
                          align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2:
                set_cell_shading(cells[i], VERY_LIGHT_BLUE)
    if widths:
        set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph(style="正文")
    if not first_indent:
        p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="要点")
        p.add_run(item)


def add_code(doc, text):
    p = doc.add_paragraph(style="代码")
    for i, line in enumerate(text.rstrip().splitlines()):
        if i:
            p.add_run().add_break()
        p.add_run(line)
    return p


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE", "1")


def style_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for name, size, color, before, after in [
        ("Title", 26, BLUE, 0, 20),
        ("Heading 1", 17, BLUE, 18, 10),
        ("Heading 2", 14, BLUE2, 14, 7),
        ("Heading 3", 12, "365F91", 10, 5),
        ("Caption", 9, GRAY, 4, 8),
    ]:
        st = doc.styles[name]
        st.font.name = "黑体" if name != "Caption" else "宋体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), st.font.name)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = name != "Caption"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "正文" not in [s.name for s in doc.styles]:
        body = doc.styles.add_style("正文", WD_STYLE_TYPE.PARAGRAPH)
    else:
        body = doc.styles["正文"]
    body.font.name = "宋体"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    body.font.size = Pt(10.5)
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body.paragraph_format.space_after = Pt(4)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    bullet = doc.styles.add_style("要点", WD_STYLE_TYPE.PARAGRAPH)
    bullet.base_style = body
    bullet.font.name = "宋体"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    bullet.paragraph_format.left_indent = Cm(0.74)
    bullet.paragraph_format.first_line_indent = Cm(-0.5)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.35

    code = doc.styles.add_style("代码", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code.font.size = Pt(8)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.right_indent = Cm(0.5)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0
    p_pr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)

    # Link heading levels to outline levels and keep heading numbers in text.
    for level in (1, 2, 3):
        st = doc.styles[f"Heading {level}"]
        ppr = st.element.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level - 1))
        ppr.append(outline)

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def load_performance():
    wb = load_workbook(PERF, data_only=True, read_only=True)
    ws = wb["Clang-Tidy"]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not values[1]:
            continue
        rows.append({
            "difficulty": values[0], "rule": values[1], "time": float(values[2]),
            "tokens": int(values[3]), "cost": float(values[4]),
            "pass_rate": float(str(values[5]).strip("%")),
            "recall": float(str(values[6]).strip("%")),
        })
    ws2 = wb["Sheet2"]
    difficulty = []
    last_diff = None
    for values in ws2.iter_rows(min_row=2, values_only=True):
        if values[0] is not None:
            last_diff = values[0]
        if values[1] == "Clang-Tidy":
            difficulty.append({
                "difficulty": last_diff, "time": float(values[3]), "tokens": int(values[4]),
                "cost": float(values[5]), "pass_rate": float(str(values[6]).strip("%")),
                "recall": float(str(values[7]).strip("%")),
            })
    return rows, difficulty


def knowledge_counts():
    with open(ROOT / "clang_tidy_collect/collect_clang_tidy_astMatchers/astMatchers.json", encoding="utf-8") as f:
        matcher_data = json.load(f)
    matchers = sum(len(next(iter(section.values()))["matchers"]) for section in matcher_data)
    with open(ROOT / "clang_tidy_collect/collect_clang_tidy_astMatchers_op/astMatchers_meta_op.json", encoding="utf-8") as f:
        matcher_ops = len(json.load(f))
    with open(ROOT / "clang_tidy_collect/collect_check_op/clang_tidy_check_op_dedup.json", encoding="utf-8") as f:
        check_ops = len(json.load(f))
    with open(ROOT / "clang_tidy_collect/collect_clang_tidy_ast_api/clang_tidy_ast_api.json", encoding="utf-8") as f:
        ast_api = json.load(f)
    api_types = sum(len(v) for v in ast_api.values())
    api_methods = sum(len(item.get("methods", [])) for v in ast_api.values() for item in v)
    return matchers, matcher_ops, check_ops, api_types, api_methods


def make_architecture(path):
    fig, ax = plt.subplots(figsize=(12, 5.5))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6)
    ax.axis("off")
    boxes = [
        (0.25, 2.15, 1.7, 1.45, "Rule +\nTest Suite", "#D9EAF7"),
        (2.35, 2.15, 1.75, 1.45, "AST & Logic\nInduction", "#B4C7E7"),
        (4.5, 2.15, 1.75, 1.45, "Knowledge\nRetrieval", "#9DC3E6"),
        (6.65, 2.15, 1.75, 1.45, "Candidate\nSynthesis", "#70AD47"),
        (8.8, 2.15, 1.45, 1.45, "Compile\nCheck", "#FFD966"),
        (10.65, 2.15, 1.1, 1.45, "Suite\nTest", "#F4B183"),
    ]
    for x, y, w, h, label, color in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.08,rounding_size=0.1",
                                    facecolor=color, edgecolor="#1F4E79", linewidth=1.5))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=11, weight="bold")
    for a, b in zip(boxes[:-1], boxes[1:]):
        x1 = a[0] + a[2]; x2 = b[0]
        ax.add_patch(FancyArrowPatch((x1 + .04, 2.88), (x2 - .04, 2.88), arrowstyle="-|>",
                                     mutation_scale=14, linewidth=1.4, color="#1F4E79"))
    ax.add_patch(FancyArrowPatch((9.55, 2.12), (7.52, 1.1), connectionstyle="arc3,rad=0.25",
                                 arrowstyle="-|>", mutation_scale=14, linewidth=1.4, color="#C00000"))
    ax.text(8.55, .72, "compiler diagnostics -> repair", ha="center", color="#C00000", fontsize=10)
    ax.add_patch(FancyArrowPatch((11.2, 2.08), (3.2, 1.25), connectionstyle="arc3,rad=0.2",
                                 arrowstyle="-|>", mutation_scale=14, linewidth=1.4, color="#7030A0"))
    ax.text(7.15, .35, "false positive / missed case -> logic augmentation", ha="center",
            color="#7030A0", fontsize=10)
    ax.text(6, 5.15, "Test-case-driven static checker generation loop", ha="center",
            fontsize=17, color="#1F4E79", weight="bold")
    ax.text(6, 4.55, "Structured evidence constrains every generation and repair step",
            ha="center", fontsize=10, color="#666666")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_feedback(path):
    fig, ax = plt.subplots(figsize=(11, 6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 7)
    ax.axis("off")
    nodes = {
        "seed": (0.5, 4.9, 2.1, .9, "Violation seed"),
        "gen": (3.35, 4.9, 2.1, .9, "Generate / update"),
        "compile": (6.25, 4.9, 1.6, .9, "Compile"),
        "suite": (6.25, 2.75, 1.6, .9, "Full suite"),
        "classify": (3.35, 2.75, 2.1, .9, "Classify failure"),
        "done": (8.5, 2.75, 1.1, .9, "Accept"),
    }
    for key, (x, y, w, h, label) in nodes.items():
        color = "#70AD47" if key == "done" else "#D9EAF7"
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=.08",
                                    facecolor=color, edgecolor="#1F4E79", linewidth=1.4))
        ax.text(x+w/2, y+h/2, label, ha="center", va="center", fontsize=11, weight="bold")
    def arrow(p1, p2, color="#1F4E79", label=None, lx=None, ly=None):
        ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=14,
                                     linewidth=1.4, color=color, connectionstyle="arc3"))
        if label:
            ax.text(lx, ly, label, fontsize=9, color=color, ha="center")
    arrow((2.6, 5.35), (3.35, 5.35))
    arrow((5.45, 5.35), (6.25, 5.35))
    arrow((7.05, 4.9), (7.05, 3.65), label="success", lx=7.48, ly=4.25)
    arrow((7.85, 3.2), (8.5, 3.2), label="all pass", lx=8.16, ly=3.45)
    arrow((6.25, 3.2), (5.45, 3.2), color="#7030A0")
    arrow((4.4, 3.65), (4.4, 4.9), color="#7030A0", label="FP / miss", lx=4.95, ly=4.28)
    ax.add_patch(FancyArrowPatch((6.55, 5.8), (5.2, 5.8), arrowstyle="-|>", mutation_scale=14,
                                 linewidth=1.4, color="#C00000", connectionstyle="arc3,rad=.35"))
    ax.text(5.95, 6.45, "compile error + retrieved repair evidence", fontsize=9,
            color="#C00000", ha="center")
    ax.text(4.4, 2.1, "Expected report -> missed detection repair\nExpected silence -> false-positive repair",
            ha="center", va="center", fontsize=10, color="#7030A0")
    ax.text(5, .75, "Bounded retries + regression testing + best-so-far state", ha="center",
            fontsize=13, color="#1F4E79", weight="bold")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_knowledge_chart(path, counts):
    matchers, matcher_ops, check_ops, _, api_methods = counts
    labels = ["Matcher API", "Matcher Meta-OP", "Check Meta-OP", "AST methods"]
    values = [matchers, matcher_ops, check_ops, api_methods]
    fig, ax = plt.subplots(figsize=(9.5, 4.8))
    bars = ax.bar(labels, values, color=["#5B9BD5", "#4472C4", "#70AD47", "#ED7D31"])
    ax.set_ylabel("Records / methods")
    ax.set_title("Knowledge-base inventory in the current repository", color="#1F4E79", weight="bold")
    ax.grid(axis="y", linestyle="--", alpha=.25)
    ax.spines[["top", "right"]].set_visible(False)
    for b, v in zip(bars, values):
        ax.text(b.get_x()+b.get_width()/2, v + max(values)*.02, f"{v:,}", ha="center", fontsize=10)
    ax.set_ylim(0, max(values)*1.14)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_performance_chart(path, difficulty):
    labels = ["Easy", "Medium", "Hard"]
    pass_rate = [x["pass_rate"] for x in difficulty]
    recall = [x["recall"] for x in difficulty]
    times = [x["time"] for x in difficulty]
    tokens = [x["tokens"] / 1000 for x in difficulty]
    x = range(3)
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11, 4.7))
    width = .34
    ax1.bar([i-width/2 for i in x], pass_rate, width, label="Test pass rate", color="#5B9BD5")
    ax1.bar([i+width/2 for i in x], recall, width, label="Recall", color="#70AD47")
    ax1.set_xticks(list(x), labels)
    ax1.set_ylim(90, 101.5)
    ax1.set_ylabel("Percent (%)")
    ax1.set_title("Correctness by difficulty", color="#1F4E79", weight="bold")
    ax1.legend(fontsize=8, loc="lower left")
    ax1.grid(axis="y", linestyle="--", alpha=.25)
    for vals, offset in ((pass_rate, -width/2), (recall, width/2)):
        for i, v in enumerate(vals): ax1.text(i+offset, v+.25, f"{v:.1f}", ha="center", fontsize=8)
    ax2.bar([i-width/2 for i in x], times, width, label="Time (s)", color="#ED7D31")
    ax2.set_xticks(list(x), labels)
    ax2.set_ylabel("Mean time (s)", color="#ED7D31")
    ax2.tick_params(axis="y", labelcolor="#ED7D31")
    ax2.set_title("Resource cost by difficulty", color="#1F4E79", weight="bold")
    ax2.grid(axis="y", linestyle="--", alpha=.2)
    ax3 = ax2.twinx()
    ax3.plot(list(x), tokens, marker="o", linewidth=2.2, color="#7030A0", label="Tokens (k)")
    ax3.set_ylabel("Mean tokens (thousand)", color="#7030A0")
    ax3.tick_params(axis="y", labelcolor="#7030A0")
    handles = [ax2.patches[0], ax3.lines[0]]
    ax2.legend(handles, ["Time (s)", "Tokens (k)"], fontsize=8, loc="upper left")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("测试用例驱动的\n静态代码检查器生成技术")
    r.bold = True
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AutoChecker 技术报告\n——面向 clang-tidy 检查器生成链路")
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    p.paragraph_format.line_spacing = 1.5

    doc.add_paragraph().paragraph_format.space_after = Pt(100)
    meta = add_table(doc, ["文档属性", "内容"], [
        ("技术对象", "AutoChecker 的 C/C++ 静态代码检查器自动生成能力"),
        ("实现载体", "clang-tidy / LLVM Clang AST"),
        ("版本依据", "当前工作区代码与 AutoChecker_Performance.xlsx"),
        ("形成日期", "2026 年 9 月"),
    ], widths=[4, 11], font_size=10)
    meta.rows[0].cells[0].merge(meta.rows[0].cells[1])
    set_cell_text(meta.rows[0].cells[0], "文档信息", bold=True, color=WHITE, size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_abstract(doc):
    doc.add_heading("摘  要", level=1)
    add_body(doc, "针对定制静态代码检查器开发依赖编译器前端知识、接口体系复杂、调试周期长等问题，本报告提出并实现一种测试用例驱动的静态代码检查器生成技术。该技术以自然语言规则和带期望结果的正反测试套件为共同规格：首先从应报告的违规样例中提取抽象语法树证据，将检测意图拆分为模式定位逻辑与语义验证逻辑；随后在匹配接口、匹配元操作、验证元操作和抽象语法树接口四类知识库中进行语义检索，并通过样例实际出现的节点类型对候选知识进行重排；在此基础上由生成模型合成候选检查器源代码。候选结果不是直接交付，而是依次通过编译验证和全量测试验证。编译诊断驱动代码级修复，误报与漏报样例驱动检测逻辑增强，从而构成“生成—验证—反馈—再生成”的双闭环。")
    add_body(doc, "当前工程以 clang-tidy 作为 C/C++ 静态代码检查器的落地载体，但方法章节按框架无关的检查器生成过程进行抽象。工程实现包含模型调用重试、结构化输出校验、AST 缓存、检索模型缓存、测试并行、轮次上限、过程快照、断点续跑和环境恢复等机制。基于 AutoChecker_Performance.xlsx 中 Clang-Tidy 明细页的 10 条规则、200 个测试用例统计，平均测试通过率为 99.5%，违规样例召回率为 99.0%；其中 9 条规则达到 100% 测试通过率，表明测试反馈闭环能够有效约束生成结果。")
    p = doc.add_paragraph(style="正文")
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("关键词：")
    r.bold = True
    p.add_run("AutoChecker；测试用例驱动；静态代码检查器；大语言模型；知识检索；闭环验证；clang-tidy")
    doc.add_page_break()


def add_toc(doc):
    doc.add_heading("目  录", level=1)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "请在 Word 中右键选择“更新域”以刷新目录和页码。")
    p.paragraph_format.line_spacing = 1.5
    doc.add_page_break()


def build_report():
    DOCS.mkdir(exist_ok=True)
    ASSETS.mkdir(exist_ok=True)
    rows, difficulty = load_performance()
    counts = knowledge_counts()
    make_architecture(ASSETS / "architecture.png")
    make_feedback(ASSETS / "feedback_loop.png")
    make_knowledge_chart(ASSETS / "knowledge_inventory.png", counts)
    make_performance_chart(ASSETS / "performance.png", difficulty)

    doc = Document()
    style_document(doc)
    add_title_page(doc)
    add_abstract(doc)
    add_toc(doc)

    # Chapter 1
    doc.add_heading("1 技术概述", level=1)
    doc.add_heading("1.1 背景与问题", level=2)
    add_body(doc, "静态代码检查将代码质量保障前移到编译和集成阶段，能够在程序运行前识别编码规范违例、潜在缺陷和安全风险。通用引擎虽然内置大量规则，但组织标准、行业标准和项目约束往往具有专用语义；当现有规则不能覆盖需求时，开发者需要理解抽象语法树、匹配接口、节点绑定、源代码位置和诊断机制，再经历多轮编译与测试，定制成本较高。")
    add_body(doc, "AutoChecker 将规则描述与测试套件视为同等重要的输入。规则文本给出意图和边界，测试用例给出可执行的行为判据。生成目标不只是“语法看起来合理”的代码，而是能够接入目标静态分析框架、通过编译，并对每个样例产生与期望一致的报告行为。")

    doc.add_heading("1.2 技术范围与术语约定", level=2)
    add_body(doc, "本报告只讨论 AutoChecker 生成 C/C++ clang-tidy checker 的部分，不覆盖仓库中的 CodeQL、Semgrep、PMD 或规则文档解析链路。为突出方法的可迁移性，第 2～3 章统一使用“静态代码检查器”“目标分析框架”等术语；第 4～6 章在说明工程实现、文件接口和实验载体时使用 clang-tidy。")
    add_table(doc, ["术语", "定义"], [
        ("规则描述", "以自然语言描述的检测对象、违规条件、允许情形和诊断意图。"),
        ("违规例（应报告）", "包含目标违例且期望检查器产生至少一条 warning 的测试用例。工程中由 CHECK-MESSAGES 标记，并映射为 case_flag=False。"),
        ("合规例（应静默）", "不包含目标违例且期望检查器不产生 warning 的测试用例。工程中映射为 case_flag=True。"),
        ("检测逻辑", "结构化的细粒度步骤，分为模式定位逻辑和命中后验证逻辑。"),
        ("知识上下文", "从四类知识库检索得到的接口说明、签名和可复用实现片段。"),
        ("候选检查器", "已合成但尚未同时通过编译与完整测试套件验证的检查器代码。"),
    ], widths=[3.2, 12.0], font_size=9)

    doc.add_heading("1.3 总体目标与设计原则", level=2)
    add_bullets(doc, [
        "• 可执行性：最终结果必须能够编译并注册到目标静态分析框架。",
        "• 行为一致性：违规例应报告、合规例应静默，以测试套件作为验收契约。",
        "• 知识约束性：生成模型使用检索到的真实接口和实现片段，降低接口幻觉。",
        "• 渐进收敛性：从单个违规样例建立可运行种子，再用失败样例逐步扩展边界。",
        "• 工程可恢复性：限制重试轮次，保留中间代码和结果检查点，并在规则结束后恢复共享构建环境。",
    ])

    doc.add_heading("1.4 总体架构", level=2)
    add_body(doc, "系统由输入建模、AST 与逻辑归纳、知识增强检索、候选代码合成、编译验证、测试验证六个环节组成。两条反馈路径分别处理“代码不可构建”和“行为不符合样例”两类问题。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "architecture.png"), width=Inches(6.6))
    add_caption(doc, "图 1  AutoChecker 测试用例驱动的静态代码检查器生成总体架构")
    add_body(doc, "与一次性生成相比，该架构把模型输出置于确定性的编译器和测试判据之下：知识检索提高首次生成质量，双闭环负责发现并修复剩余偏差，全量回归测试防止局部修复破坏已通过样例。")

    # Chapter 2
    doc.add_heading("2 测试用例驱动的输入建模", level=1)
    doc.add_heading("2.1 规则—样例联合规格", level=2)
    add_body(doc, "将单条生成任务形式化为 T=(r, d, C)，其中 r 为规则标识，d 为自然语言规则描述，C={(xᵢ,yᵢ)} 为测试套件；xᵢ 是源代码，yᵢ∈{report,silent} 表示期望行为。规则文本通常不能穷举所有语法形态，而样例能够把“应该报什么、不应该报什么”转化为可运行断言。二者联合可同时约束检测目标和排除边界。")
    add_table(doc, ["输入字段", "当前实现", "生成阶段用途"], [
        ("main_title", "规则名称", "生成类名、源文件名、结果目录及框架注册名。"),
        ("description", "规则语义与合规/违规场景", "逻辑归纳、代码生成和反馈增强。"),
        ("rule_test_path", "测试目录", "加载目录下全部 .cpp 用例并执行回归。"),
        ("category", "检查器模块类别", "确定测试包含目录与框架模块。"),
        ("CHECK-MESSAGES", "样例内期望诊断标记", "区分应报告与应静默样例，并由测试工具核验。"),
    ], widths=[3.1, 4.3, 7.8], font_size=8.8)

    doc.add_heading("2.2 样例语义与判定函数", level=2)
    add_body(doc, "当前代码沿用测试框架的命名习惯：包含 CHECK-MESSAGES 的用例被称为“negative case”，表示代码违反规则且应产生警告；不包含该标记的用例被称为“positive case”，表示代码合规且应保持静默。为避免语义歧义，本报告分别称为“违规例”和“合规例”。")
    add_code(doc, "Pass(x_i) = { warning_count(x_i) > 0,  if y_i = report\n              { warning_count(x_i) = 0,  if y_i = silent\n\nSuitePass = AND_i Pass(x_i)")
    add_body(doc, "测试执行返回负值时表示工具运行错误，不计为满足期望。当前验证以 warning 数量为判据，违规例只要求至少命中一次；因此报告位置和消息内容的更细粒度一致性主要交由 CHECK-MESSAGES 所依赖的测试脚本完成。")

    doc.add_heading("2.3 基于抽象语法树的样例表征", level=2)
    doc.add_heading("2.3.1 抽象语法树提取与裁剪", level=3)
    add_body(doc, "系统调用目标编译器前端，以“-Xclang -ast-dump -fsyntax-only”获取样例的文本 AST。为减少系统头文件和前置声明造成的噪声，处理逻辑保留首行，并从首个以“`-”开头的根分支起截取后续内容。该表示连同原始代码一起进入生成提示，使模型能够把自然语言规则映射到实际节点层级。")
    doc.add_heading("2.3.2 节点类型集合与缓存", level=3)
    add_body(doc, "系统利用“节点类型 + 内存地址”模式解析 AST，按地址去重，再形成节点类型集合。节点集合不仅用于提示模型，还用于后续检索结果的相关性重排。相同样例的 AST 结果按绝对路径缓存在进程内，避免增强轮次反复执行前端分析。")
    add_table(doc, ["处理步骤", "输入", "输出", "作用"], [
        ("前端解析", ".cpp 文件", "AST 文本", "获得与编译器一致的语法结构。"),
        ("树裁剪", "完整 AST", "目标根分支", "降低系统声明噪声和提示长度。"),
        ("节点解析", "裁剪后的 AST", "节点类型集合", "为知识检索提供结构先验。"),
        ("路径缓存", "样例绝对路径", "AST 三元组", "减少重复前端调用。"),
    ], widths=[2.8, 3.0, 3.5, 5.9], font_size=8.8)

    doc.add_heading("2.4 种子样例选择策略", level=2)
    add_body(doc, "初始检查器必须先证明能够识别至少一个目标违例，因此系统按测试集合顺序选择尚未跳过的违规例作为种子。若当前种子在最大生成轮次内未获得“可编译且能报告”的检查器，则将其加入跳过集合并尝试下一违规例。得到种子检查器后，系统清空临时跳过状态并进入完整测试集增强阶段，使初始选择只影响启动路径，不改变最终验收范围。")

    # Chapter 3
    doc.add_heading("3 静态代码检查器生成方法", level=1)
    doc.add_heading("3.1 样例约束的检测逻辑归纳", level=2)
    doc.add_heading("3.1.1 规则语义与违规模式联合归纳", level=3)
    add_body(doc, "生成模型接收规则描述和当前违规例，不直接输出完整源码，而是先解释“目标节点如何定位、命中节点如何验证”。测试样例提供具体语法模式，规则描述用于防止模型把单一样例的变量名、字面量或局部结构误当作规则本身。")
    doc.add_heading("3.1.2 双通道检测逻辑中间表示", level=3)
    add_body(doc, "当前实现将逻辑表示为一个非空 JSON 数组，其中首个对象必须包含 logic_registerMatchers 与 logic_check 两个非空字符串数组。前者描述候选节点的结构定位与绑定，后者描述命中后的属性读取、上下文判断、排除条件和诊断输出。该中间表示把框架中的“广域模式筛选”和“精细语义确认”分离，便于检索不同类型的知识。")
    add_code(doc, '[\n  {\n    "logic_registerMatchers": [\n      "定位控制流条件中的赋值运算节点",\n      "绑定赋值节点与所在条件语句"\n    ],\n    "logic_check": [\n      "读取绑定的赋值运算符位置",\n      "排除允许的比较表达式上下文",\n      "在有效源码位置报告诊断"\n    ]\n  }\n]')
    doc.add_heading("3.1.3 结构校验与模型重试", level=3)
    add_body(doc, "模型输出首先移除 Markdown 代码围栏，再进行 JSON 解析和字段级校验。空响应、非数组、空逻辑列表或非字符串步骤均被判为无效；系统最多重试 max_llm_tries 次，当前配置为 3。该机制将开放式模型输出转换为稳定的程序接口，防止错误在检索和代码生成阶段放大。")

    doc.add_heading("3.2 检测逻辑引导的知识增强检索", level=2)
    doc.add_heading("3.2.1 四类知识库协同", level=3)
    add_body(doc, "模式定位逻辑分别检索“匹配接口库”和“匹配元操作库”，语义验证逻辑分别检索“AST 接口库”和“验证元操作库”。接口库提供名称、签名、参数和说明，元操作库提供从既有检查器中抽取的“语义描述—实现片段”对。两者结合兼顾 API 正确性和组合用法。")
    matchers, matcher_ops, check_ops, api_types, api_methods = counts
    add_table(doc, ["知识库", "当前规模", "检索内容", "主要用途"], [
        ("匹配接口库", f"{matchers:,} 个匹配器", "节点、收窄、遍历匹配器的签名与说明", "构造候选节点模式。"),
        ("匹配元操作库", f"{matcher_ops:,} 条", "匹配语义与实现片段", "复用成熟的匹配组合。"),
        ("验证元操作库", f"{check_ops:,} 条", "节点获取、条件判断与诊断片段", "完成命中后的验证。"),
        ("AST 接口库", f"{api_types:,} 个类型、{api_methods:,} 个方法", "类/结构体方法签名", "读取节点属性与上下文。"),
    ], widths=[3.0, 3.1, 5.7, 3.7], font_size=8.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "knowledge_inventory.png"), width=Inches(6.25))
    add_caption(doc, "图 2  当前仓库中四类生成知识的规模（由 JSON 数据文件统计）")

    doc.add_heading("3.2.2 语义向量检索", level=3)
    add_body(doc, "系统使用本地 BGE 大文本嵌入模型对逻辑步骤和知识条目编码，向量归一化后以内积计算余弦相似度。每个逻辑步骤取相似度最高的 k 条记录，当前 k=5；多步骤结果合并后去重。知识向量持久化为 .pt 文件，模型采用线程安全的进程内单例缓存，减少重复加载约 1.3 GB 模型的开销。")
    add_code(doc, "score(q, d) = normalize(E(q)) · normalize(E(d))\nR(q) = TopK_d score(q, d),  k = 5")
    doc.add_heading("3.2.3 AST 感知的候选重排", level=3)
    add_body(doc, "仅依赖语义相似度可能召回含义接近但节点类型不适用于当前样例的接口。系统因此扫描检索文档中出现的 AST 类型：与样例节点类型精确匹配加 1.0 分，去除 CXX 前缀后匹配加 0.8 分，通用遍历匹配器加 0.3 分。候选按该分数重排并截取前 3 条；若所有分数为 0，则回退到原始语义排序。该策略以样例结构作为硬证据，同时避免对未知模式进行过度过滤。")

    doc.add_heading("3.3 结构化候选代码合成", level=2)
    doc.add_heading("3.3.1 框架模板初始化", level=3)
    add_body(doc, "系统先调用框架提供的新增检查器脚本生成类定义、源文件、构建项和模块注册骨架。生成模型在既有骨架上补全，而不是自行猜测工程接入方式。处理每条规则前先检查并清理中断遗留模板，预编译基础工程，确保后续错误主要来自本次生成内容。")
    doc.add_heading("3.3.2 多源上下文组装", level=3)
    add_body(doc, "候选生成提示由规则描述、种子代码、裁剪后的 AST、双通道检测逻辑、匹配类知识、验证类知识，以及当前 .cpp/.h 模板共同组成。增强阶段还加入当前检查器、已通过样例和当前失败样例。这样，生成模型同时获得“要检测什么”“样例结构是什么”“可调用什么”“已有行为不能破坏”四类约束。")
    doc.add_heading("3.3.3 双文件输出解析与落盘", level=3)
    add_body(doc, "模型必须输出实现文件和头文件。解析器优先识别 checker_cpp/checker_h 标签后的 cpp 或 c++ 代码块；若标签缺失，则回退到前两个 C++ 代码块。只有两个代码块均非空时才进入后续阶段。候选代码写入框架工作树，同时按“初始生成—轮次—编译修复—最终版本”保存过程快照，便于复现和问题定位。")

    doc.add_heading("3.4 编译反馈驱动的代码修复", level=2)
    doc.add_heading("3.4.1 编译诊断清洗", level=3)
    add_body(doc, "原始构建输出包含进度、命令行和 CMake 噪声。系统移除常规构建进度、完整编译命令、FAILED 标记等非诊断内容，保留语法错误、类型错误、符号缺失和模板实例化信息；若清洗结果为空则回退到末尾 3000 字符，并将最大行数限制为 500。清洗后的上下文更聚焦，也降低模型输入成本。")
    doc.add_heading("3.4.2 错误分析—知识补检索—修复", level=3)
    add_body(doc, "错误分析模型接收当前源代码、头文件和清洗后的编译信息，输出 repair_step 与 wait_retrieve_code_snippet。后者作为提示再次检索四类知识库，补充与错误相关的接口或实现片段；随后代码生成模型依据修复步骤、检索建议和原代码生成完整替换版本。当前每个生成轮次最多进行 2 次编译修复。")

    doc.add_heading("3.5 测试反馈驱动的行为增强", level=2)
    doc.add_heading("3.5.1 全量测试与失败分类", level=3)
    add_body(doc, "种子检查器通过单个违规例后，系统并行运行全部未跳过用例。合规例出现警告被归为误报，违规例未出现警告被归为漏报；失败样例连同期望语义进入不同的增强提示。当前默认使用 4 个工作线程，每个用例拥有独立临时文件和日志文件，避免并发写冲突。")
    doc.add_heading("3.5.2 面向误报与漏报的差异化增强", level=3)
    add_body(doc, "对于漏报，逻辑增强侧重补充语法变体、上下文路径和目标节点覆盖；对于误报，逻辑增强侧重增加排除条件、作用域限制和语义守卫。系统基于增强逻辑重新检索知识，再生成完整检查器。当前失败样例修复成功后立即重新运行全量套件，只有不破坏既有通过集合的版本才会作为新的候选状态。")
    doc.add_heading("3.5.3 有界迭代与终止条件", level=3)
    add_body(doc, "对单个失败样例，生成增强轮次受 max_round 限制，当前为 2；每次模型格式失败还受 max_llm_tries=3 限制。达到上限的样例进入跳过集合，以保证批处理可终止。流程在全部有效样例通过时成功结束；若初始阶段所有违规种子均失败，则该规则生成失败。")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "feedback_loop.png"), width=Inches(6.35))
    add_caption(doc, "图 3  编译反馈与测试反馈构成的双闭环")

    doc.add_heading("3.6 核心算法", level=2)
    add_code(doc, "输入：规则描述 d，测试套件 C，最大轮次 R，编译修复上限 B\n输出：通过验证的静态代码检查器或失败状态\n\n1  初始化框架模板并预编译基础工程\n2  for 每个未尝试的违规种子 c in C:\n3      提取并缓存 AST(c)，归纳双通道检测逻辑 L\n4      K <- 语义检索(L) + AST 感知重排\n5      q <- 基于 d、c、AST(c)、L、K 合成候选检查器\n6      在 B 次上限内：编译 q；失败则分析诊断、补检索并修复 q\n7      若 q 可编译且能命中 c，则进入步骤 9；否则换种子\n8  若不存在有效种子，返回失败\n9  循环执行完整测试套件，得到失败集合 F\n10 若 F 为空，返回 q\n11 对 F 中样例 f：按误报/漏报生成增强逻辑，检索知识并更新 q\n12 对更新后的 q 重复编译修复，并回归完整测试套件\n13 超过轮次上限的样例加入跳过集合；按终止条件返回结果")

    # Chapter 4
    doc.add_heading("4 AutoChecker 工程实现", level=1)
    doc.add_heading("4.1 模块划分与职责", level=2)
    add_table(doc, ["模块", "关键文件", "职责"], [
        ("批处理编排", "src/main.py", "加载规则、创建模板、调用生成器、保存结果、断点续跑、清理和恢复环境。"),
        ("生成与反馈", "src/generator.py", "逻辑归纳、候选生成、编译修复、测试增强、用量统计。"),
        ("AST 与检索协调", "src/help/clang_tidy_utils.py", "AST 提取/缓存、逻辑格式化、四库检索、AST 重排、代码块解析。"),
        ("构建与测试适配", "src/plateform/clang_tidy.py", "模板增删、clang-tidy 构建、测试脚本执行、诊断清洗。"),
        ("知识检索器", "src/retriever/*.py", "BGE 编码、向量缓存、Top-k 相似度检索、元操作到实现片段映射。"),
        ("提示模板", "src/prompt/clang_tidy_prompt/", "按阶段维护系统提示与用户提示，隔离提示工程和流程代码。"),
        ("领域实体", "src/entity/", "封装规则、样例、检查器及已通过样例状态。"),
    ], widths=[2.6, 4.8, 7.6], font_size=8.3)

    doc.add_heading("4.2 关键实现参数", level=2)
    add_table(doc, ["参数", "当前值", "实现含义"], [
        ("max_round", "2", "单个种子或失败样例的最大生成/增强轮次。"),
        ("max_compiler_trys", "2", "每轮候选的最大编译修复次数。"),
        ("max_llm_tries", "3", "结构化逻辑或双文件代码输出解析失败后的最大模型重试次数。"),
        ("top_key", "5", "每个逻辑查询在各知识库中的语义候选数。"),
        ("AST 重排截断", "3", "根据样例节点相关度重新排序后保留的候选数。"),
        ("测试并行度", "4", "runAllTestCase 默认工作线程数。"),
        ("C++ 标准", "C++17", "测试脚本默认编译参数。"),
        ("构建并行度", "56", "clang-tidy 目标的 CMake 构建并行参数。"),
    ], widths=[3.5, 2.2, 9.3], font_size=8.8)
    add_body(doc, "上述参数是当前代码的工程默认值，不是方法本身的固定常数。实际部署应结合模型稳定性、机器并行能力、构建耗时和规则复杂度调优。")

    doc.add_heading("4.3 状态管理与可追溯性", level=2)
    add_bullets(doc, [
        "• 中间快照：每个违规种子、生成轮次、首次生成和编译修复均保留 .cpp/.h 文件。",
        "• 调试提示：各阶段实际发送给模型的提示保存到 debug_prompt 目录。",
        "• 细粒度用量：按调用记录模型、阶段标签、输入/输出/缓存 token 和成本，并汇总到规则结果。",
        "• 原子写入：结果先写临时文件，再通过 os.replace 替换，避免中断留下半个 JSON。",
        "• 断点续跑：按 rule_id 加载已有结果，可跳过成功项或显式重试失败项，并保持源规则顺序。",
    ])

    doc.add_heading("4.4 共享构建环境保护", level=2)
    add_body(doc, "clang-tidy 检查器生成会修改共享 LLVM 工作树。批处理在规则开始前清理残留模板并预编译，结束时无论成功或异常都删除本次模板并重新构建。如果环境恢复失败，系统在保存当前检查点后停止批处理，避免一个规则的残留改动污染后续全部规则。这一机制是将生成实验扩展为稳定批处理的关键。")

    doc.add_heading("4.5 并发与性能优化", level=2)
    add_body(doc, "测试用例彼此独立，系统通过线程池并行执行；AST 以路径为键缓存；嵌入模型采用线程安全单例；知识向量和 API 字典持久化缓存。与此同时，知识编码保持单进程批量执行，避免多进程各自加载大模型带来的内存放大和父进程线程锁问题。优化点覆盖了前端解析、模型加载、向量计算和回归测试四个主要耗时来源。")

    # Chapter 5
    doc.add_heading("5 实验数据与效果分析", level=1)
    doc.add_heading("5.1 数据来源与统计口径", level=2)
    add_body(doc, "本章数据来自仓库根目录 AutoChecker_Performance.xlsx 的“Clang-Tidy”明细页和按难度汇总页。为保证可复核性，只采用明细页可逐条对应的 10 条 C/C++ 规则，每条规则包含 20 个测试用例（10 个违规例、10 个合规例），共 200 个样例。该工作簿是已有实验记录，本报告未重新运行完整生成实验。")
    add_table(doc, ["指标", "定义/口径"], [
        ("测试通过率", "每条规则中行为符合期望的样例数 / 20，再对规则取平均。"),
        ("违规召回率", "被检查器成功报告的违规例数 / 10，再对规则取平均。"),
        ("耗时", "单条规则完成生成、修复和测试流程的记录时间，单位秒。"),
        ("Token 消耗", "单条规则各模型调用的输入与输出 token 记录汇总。"),
        ("花费", "工作簿按调用时模型价格计算的单规则成本，单位人民币元。"),
    ], widths=[3.2, 11.8], font_size=8.8)

    doc.add_heading("5.2 总体效果", level=2)
    mean_time = sum(x["time"] for x in rows)/len(rows)
    mean_tokens = sum(x["tokens"] for x in rows)/len(rows)
    mean_cost = sum(x["cost"] for x in rows)/len(rows)
    mean_pass = sum(x["pass_rate"] for x in rows)/len(rows)
    mean_recall = sum(x["recall"] for x in rows)/len(rows)
    perfect = sum(x["pass_rate"] == 100 for x in rows)
    add_table(doc, ["规则数", "测试数", "平均耗时", "平均 Token", "平均成本", "测试通过率", "违规召回率"], [[
        len(rows), len(rows)*20, f"{mean_time:.2f} s", f"{mean_tokens:,.0f}", f"¥{mean_cost:.4f}", f"{mean_pass:.1f}%", f"{mean_recall:.1f}%"
    ]], widths=[1.8, 1.8, 2.2, 2.3, 2.0, 2.3, 2.3], font_size=8.5)
    add_body(doc, f"10 条规则中有 {perfect} 条达到 100% 测试通过率。按 200 个样例折算，平均 99.5% 对应 199/200 个样例行为正确；99.0% 的违规召回率对应 100 个违规例中命中 99 个。唯一未满分规则为 no-same-name-as-global-variable，其测试通过率为 95%、召回率为 90%，说明剩余误差来自一个违规变体漏检，而非合规样例误报。")

    doc.add_heading("5.3 分难度效果与成本", level=2)
    diff_rows = [(x["difficulty"], f"{x['time']:.2f}", f"{x['tokens']:,}", f"¥{x['cost']:.4f}",
                  f"{x['pass_rate']:.1f}%", f"{x['recall']:.1f}%") for x in difficulty]
    add_table(doc, ["难度", "平均耗时(s)", "平均 Token", "平均成本", "测试通过率", "违规召回率"], diff_rows,
              widths=[2.0, 2.5, 2.7, 2.2, 2.7, 2.7], font_size=8.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "performance.png"), width=Inches(6.5))
    add_caption(doc, "图 4  不同规则难度下的正确性与资源消耗")
    add_body(doc, "简单规则平均测试通过率和召回率均为 100%。中等规则平均测试通过率为 98.8%、召回率为 97.5%，主要受同名变量规则的语义边界影响。困难规则仍全部达到 100%，但平均耗时由简单规则的 994.02 秒上升至 2226.11 秒，平均 token 由 111,937 增长至 271,596，表明复杂上下文和多轮反馈主要增加生成成本，而验证闭环仍能维持结果正确性。")

    doc.add_heading("5.4 逐规则结果", level=2)
    detail = []
    dmap = {"简": "简单", "中": "中等", "困": "困难"}
    for x in rows:
        detail.append((dmap.get(x["difficulty"], x["difficulty"]), x["rule"], f"{x['time']:.2f}",
                       f"{x['tokens']:,}", f"¥{x['cost']:.4f}", f"{x['pass_rate']:.1f}%", f"{x['recall']:.1f}%"))
    add_table(doc, ["难度", "规则", "耗时(s)", "Token", "成本", "通过率", "召回率"], detail,
              widths=[1.4, 5.3, 1.8, 2.1, 1.7, 1.7, 1.7], font_size=7.5)

    doc.add_heading("5.5 结果解释与局限", level=2)
    add_bullets(doc, [
        "• 数据证明的是给定测试套件内的行为一致性，不能等同于对所有真实项目代码的完备正确性。",
        "• 测试集每条规则固定为 10 个违规例和 10 个合规例，类别均衡有利于解释指标，但不代表生产环境中的违例分布。",
        "• 当前详细样本量为 10 条规则；知识库规模较大，但跨项目、跨编译选项和宏配置的泛化能力仍需扩大实验验证。",
        "• 平均值受多轮大模型推理和 LLVM 全量构建影响；困难规则的时间/token 明显上升，后续可通过增量构建、候选补丁生成和自适应预算降低成本。",
        "• “达到测试全通过”可能出现对样例过拟合，应加入独立留出集、变异测试和真实项目扫描验证。",
    ])

    # Chapter 6
    doc.add_heading("6 典型案例：条件表达式中的赋值", level=1)
    doc.add_heading("6.1 规则与测试契约", level=2)
    add_body(doc, "以 no-assignment-in-condition 为例，规则禁止在 if、while、for 等逻辑条件中直接使用赋值操作，以降低将“=”误写为“==”所导致的逻辑缺陷。该规则的测试套件包含 10 个违规例与 10 个合规例。违规例通过 CHECK-MESSAGES 指定期望诊断，合规例不包含期望警告。")
    add_code(doc, "// 违规例：应报告\nint a = 0, b = 5;\nif (a = b) {\n  // CHECK-MESSAGES: 禁止将赋值语句作为逻辑表达式 [...]\n}\n\n// 合规例：应静默\nint a = 5, b = 5;\nif (a == b) { /* ... */ }")

    doc.add_heading("6.2 生成逻辑与实现结构", level=2)
    add_body(doc, "生成结果以赋值二元运算节点为核心，排除位于比较运算父节点中的允许模式；随后覆盖逻辑与/或、if、while、do-while、for、条件运算符等上下文并绑定赋值节点。在验证阶段读取绑定节点，确认源码位置有效，识别所在条件上下文并输出诊断。该结构体现了“模式定位负责召回、命中验证负责精度”的双通道逻辑。")
    add_table(doc, ["逻辑层", "案例中的实现意图"], [
        ("目标定位", "定位 isAssignmentOperator() 对应的 BinaryOperator。"),
        ("排除约束", "排除其父节点为比较运算符的情形。"),
        ("上下文覆盖", "覆盖控制流条件、逻辑运算和条件运算符。"),
        ("节点绑定", "以 assign_in_cond 绑定待诊断的赋值节点。"),
        ("有效性检查", "节点不存在或运算符位置无效时直接返回。"),
        ("诊断输出", "在赋值运算符位置报告规则消息，并附上下文说明。"),
    ], widths=[3.2, 11.8], font_size=8.8)

    doc.add_heading("6.3 验证结果", level=2)
    x = next(r for r in rows if r["rule"] == "no-assignment-in-condition")
    add_body(doc, f"工作簿记录显示，该规则最终通过 20/20 个测试用例，测试通过率与违规召回率均为 100%；生成耗时 {x['time']:.2f} 秒，Token 消耗 {x['tokens']:,}，记录成本 ¥{x['cost']:.4f}。结果目录同时保留最终 .cpp/.h、初始生成版本、轮次快照、调试提示和 JSON 结果，可用于复核从样例到检查器的全过程。")

    # Chapter 7
    doc.add_heading("7 技术总结与演进方向", level=1)
    doc.add_heading("7.1 技术特点", level=2)
    add_body(doc, "测试用例驱动的核心价值在于把静态代码检查器生成从开放式代码创作转化为可执行规格约束下的搜索与修复过程。AutoChecker 通过双通道逻辑中间表示降低任务复杂度，通过四类知识库缓解接口幻觉，通过 AST 感知重排提高知识与样例结构的一致性，再以编译器和测试套件提供确定性反馈。由此形成知识增强、结构约束和行为验证相互配合的生成闭环。")
    doc.add_heading("7.2 后续演进方向", level=2)
    add_bullets(doc, [
        "• 引入独立留出集和自动变异测试，检测对现有样例的过拟合。",
        "• 将 warning 数量判定扩展为对诊断位置、消息、修复建议和重复报告的结构化比较。",
        "• 建立“规则语义—AST 模式—API/元操作—测试结果”可追踪图，支持失败根因定位。",
        "• 根据编译错误类别、失败样例数量和历史收敛速度动态分配生成轮次与 token 预算。",
        "• 使用增量构建和最小补丁式修复，减少困难规则的全量编译时间与完整代码重生成成本。",
        "• 扩充跨项目、宏展开、模板实例化和不同语言标准下的测试，验证真实环境泛化能力。",
    ])

    doc.add_heading("附录 A  实现依据与数据追溯", level=1)
    add_table(doc, ["报告内容", "仓库依据"], [
        ("批处理、检查点、环境恢复", "src/main.py"),
        ("初始生成、编译修复、测试增强、用量统计", "src/generator.py"),
        ("AST 提取、AST 重排、四库协调、代码解析", "src/help/clang_tidy_utils.py"),
        ("模板、构建、测试、诊断清洗", "src/plateform/clang_tidy.py"),
        ("向量编码与 Top-k", "src/retriever/bge_embedding.py 及四个 retrieve_from_*.py"),
        ("知识库规模", "clang_tidy_collect/ 下四个 JSON 数据集"),
        ("实验指标", "AutoChecker_Performance.xlsx：Clang-Tidy、Sheet2"),
        ("案例检查器", "experiment/gjb8114/result/clang-tidy/no-assignment-in-condition/"),
        ("参考结构与技术表述", "docs/一种基于大语言模型的静态代码检查器生成方法v2.docx"),
    ], widths=[5.0, 10.0], font_size=8.5)
    add_body(doc, "说明：本报告中的知识库数量由当前 JSON 文件现场统计；配置参数来自 src/config.json 与函数默认参数；实验数据为工作簿已有记录。报告未把当前仓库中尚未完成的 130 条 GJB 规则测试集规模作为生成成功率证据。")

    # Document metadata and header.
    doc.core_properties.title = "测试用例驱动的静态代码检查器生成技术"
    doc.core_properties.subject = "AutoChecker clang-tidy checker generation technical report"
    doc.core_properties.author = "AutoChecker 项目组"
    doc.core_properties.keywords = "AutoChecker, 静态代码检查器, 测试用例驱动, clang-tidy"
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run("AutoChecker 技术报告  |  测试用例驱动的静态代码检查器生成技术")
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("808080")

    doc.save(OUT)
    return {
        "output": str(OUT), "rules": len(rows), "tests": len(rows)*20,
        "mean_pass": mean_pass, "mean_recall": mean_recall, "knowledge": counts,
    }


if __name__ == "__main__":
    print(json.dumps(build_report(), ensure_ascii=False, indent=2))
