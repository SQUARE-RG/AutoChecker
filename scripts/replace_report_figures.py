#!/usr/bin/env python3
"""Replace all three report figures with their readability-improved versions."""

import os
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一二章修改版.docx"
FIGURES = {
    "图1": ROOT / "docs/专利图v6-技术报告版-v2.png",
    "图2": ROOT / "docs/autochecker_chapter3_assets/图2_静态分析知识库构建方法.png",
    "图3": ROOT / "docs/autochecker_chapter3_assets/图3_检测逻辑驱动的分层检索方法.png",
}


def replace_preceding_figure(doc, caption_prefix, image_path):
    caption_index = next(
        i for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip().startswith(caption_prefix)
    )
    figure_paragraph = next(
        paragraph for paragraph in reversed(doc.paragraphs[:caption_index])
        if paragraph._p.xpath(".//w:drawing")
    )
    figure_paragraph.clear()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.add_run().add_picture(str(image_path), width=Cm(15.70))


def main():
    if not DOCUMENT.exists():
        raise FileNotFoundError(DOCUMENT)
    for image_path in FIGURES.values():
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    doc = Document(DOCUMENT)
    for caption_prefix, image_path in FIGURES.items():
        replace_preceding_figure(doc, caption_prefix, image_path)

    fd, temp_name = tempfile.mkstemp(
        prefix=".autochecker-figures-", suffix=".docx", dir=str(DOCUMENT.parent)
    )
    os.close(fd)
    try:
        doc.save(temp_name)
        os.replace(temp_name, DOCUMENT)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)

    print(DOCUMENT)


if __name__ == "__main__":
    main()
