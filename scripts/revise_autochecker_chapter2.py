#!/usr/bin/env python3
"""Revise Chapter 2 wording for an AutoChecker midterm project report."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一章修改版.docx"
OUTPUT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一二章修改版.docx"


REVISED_PARAGRAPHS = [
    "在软件研制过程中，静态代码检查用于在程序运行前分析源代码，识别编码规范违例、潜在缺陷和安全风险。与运行测试相比，静态检查能够覆盖未被执行的代码路径，并可在编码、代码评审和持续集成阶段重复使用。对于执行行业编码标准或组织内部规范的项目，静态检查结果还可以作为代码质量审查和规范符合性检查的依据。",
    "现有静态分析工具通常提供一组通用检查规则，可以覆盖常见的语言误用和代码质量问题，但难以直接满足所有项目的专用要求。一方面，不同行业标准和项目规范对同一类代码结构可能采用不同的判定条件；另一方面，具体规则往往包含适用范围、上下文限制和允许情形。内置规则无法覆盖这些要求时，需要在现有静态分析框架上增加定制检查器。",
    "定制检查器的开发需要把自然语言规则转换为可执行的程序分析逻辑。开发人员需要确定目标抽象语法树节点、节点之间的结构关系、命中后的属性判断和诊断位置，并正确使用静态分析框架提供的匹配、遍历和节点操作接口。候选实现还要经过工程编译和测试用例验证，出现接口错误、漏报或误报后需要继续修改。上述过程同时依赖规则理解、编译器前端知识和工程调试经验，是当前检查器定制工作的主要成本来源。",
    "大语言模型具备规则理解和代码生成能力，为自动生成检查器提供了基础，但直接根据规则文本生成完整代码仍存在不稳定性：规则描述难以列出全部代码形式和排除条件，模型也可能使用不适合当前节点类型的接口。针对这些问题，AutoChecker将规则描述与测试用例共同作为生成依据，并引入静态分析知识检索、编译验证和测试反馈。规则描述用于确定检查目标，违规用例和合规用例用于限定检测边界，编译错误及测试失败结果用于推动后续修正。本项目在中期阶段重点完成这条生成与验证链路，为后续扩大规则覆盖范围提供基础。",
]


def format_heading(paragraph):
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        run.font.name = "黑体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(31, 78, 121)


def format_body(paragraph):
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def build():
    doc = Document(SOURCE)
    chapter2_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "2 背景")
    chapter3_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("3 测试用例驱动"))
    if chapter3_index - chapter2_index - 1 != 4:
        raise RuntimeError("第二章正文段落数量与预期不一致，已停止修改")

    heading = doc.paragraphs[chapter2_index]
    heading.text = "2 背景"
    format_heading(heading)

    for paragraph, replacement in zip(doc.paragraphs[chapter2_index + 1:chapter3_index], REVISED_PARAGRAPHS):
        paragraph.text = replacement
        format_body(paragraph)

    doc.core_properties.subject = "项目中期检查材料（第一、二章修改版）"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
