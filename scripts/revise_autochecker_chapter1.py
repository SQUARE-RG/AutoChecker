#!/usr/bin/env python3
"""Revise only Chapter 1 of the user's AutoChecker midterm report draft."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/测试用例驱动的静态代码检查器生成技术.docx"
OUTPUT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一章修改版.docx"
ASSET_DIR = ROOT / "docs/autochecker_chapter1_assets"
FIGURE = ASSET_DIR / "图1_AutoChecker中期技术路线.png"
FONT_PATH = Path("/tmp/NotoSansCJKsc-Regular.otf")


BLUE = "#1F4E79"
MID_BLUE = "#5B9BD5"
LIGHT_BLUE = "#D9EAF7"
GREEN = "#70AD47"
LIGHT_GREEN = "#E2F0D9"
ORANGE = "#ED7D31"
LIGHT_ORANGE = "#FCE4D6"
PURPLE = "#7030A0"
GRAY = "#666666"
LIGHT_GRAY = "#F2F2F2"
RED = "#C00000"


def font(size, bold=False):
    path = FONT_PATH
    if not path.exists():
        raise FileNotFoundError(f"缺少中文字体：{path}")
    # Noto CJK Regular is sufficiently legible; emulate emphasis with stroke.
    return ImageFont.truetype(str(path), size=size)


def rounded_box(draw, xy, fill, outline=BLUE, width=4, radius=22):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered(draw, xy, text, ft, fill="#222222", spacing=8, stroke=0):
    x1, y1, x2, y2 = xy
    box = draw.multiline_textbbox((0, 0), text, font=ft, spacing=spacing, align="center", stroke_width=stroke)
    w = box[2] - box[0]
    h = box[3] - box[1]
    draw.multiline_text(((x1+x2-w)/2, (y1+y2-h)/2), text, font=ft, fill=fill,
                        spacing=spacing, align="center", stroke_width=stroke)


def arrow(draw, start, end, fill=BLUE, width=6, head=18):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2-x1) >= abs(y2-y1):
        sign = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2-sign*head, y2-head//2), (x2-sign*head, y2+head//2)]
    else:
        sign = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2-head//2, y2-sign*head), (x2+head//2, y2-sign*head)]
    draw.polygon(pts, fill=fill)


def poly_arrow(draw, points, fill=PURPLE, width=5, head=18):
    draw.line(points, fill=fill, width=width, joint="curve")
    arrow(draw, points[-2], points[-1], fill=fill, width=width, head=head)


def make_figure():
    ASSET_DIR.mkdir(exist_ok=True)
    img = Image.new("RGB", (2600, 1180), "white")
    draw = ImageDraw.Draw(img)
    ft_title = font(45)
    ft_stage = font(35)
    ft_body = font(29)
    ft_small = font(25)
    ft_note = font(23)

    centered(draw, (0, 15, 2600, 95), "AutoChecker测试用例驱动的静态代码检查器生成技术路线",
             ft_title, BLUE, stroke=1)

    # Input block.
    rounded_box(draw, (50, 235, 330, 665), LIGHT_GRAY)
    centered(draw, (50, 255, 330, 330), "输入", ft_stage, BLUE, stroke=1)
    rounded_box(draw, (85, 355, 295, 455), "white", outline="#888888", width=3, radius=14)
    centered(draw, (85, 355, 295, 455), "规则描述", ft_body)
    rounded_box(draw, (85, 490, 295, 615), "white", outline="#888888", width=3, radius=14)
    centered(draw, (85, 490, 295, 615), "测试用例\n违规 / 合规", ft_small)

    # Three online stages.
    stages = [
        (420, 165, 1040, 745, LIGHT_BLUE, "阶段1  代码模式分析与检测逻辑生成"),
        (1110, 165, 1725, 745, "#DDEBF7", "阶段2  检测逻辑引导的知识检索"),
        (1795, 165, 2545, 745, LIGHT_GREEN, "阶段3  候选检查器生成与反馈验证"),
    ]
    for x1, y1, x2, y2, color, title in stages:
        rounded_box(draw, (x1, y1, x2, y2), color, width=4, radius=24)
        centered(draw, (x1+15, y1+22, x2-15, y1+95), title, ft_stage, BLUE, stroke=1)

    # Stage 1 details.
    rounded_box(draw, (470, 310, 760, 455), "white", outline=MID_BLUE, width=3, radius=14)
    centered(draw, (470, 310, 760, 455), "AST提取与\n代码模式分析", ft_body)
    rounded_box(draw, (790, 310, 990, 455), "white", outline=MID_BLUE, width=3, radius=14)
    centered(draw, (790, 310, 990, 455), "检测逻辑\n生成", ft_body)
    arrow(draw, (760, 383), (790, 383), MID_BLUE, 5, 15)
    rounded_box(draw, (515, 535, 945, 675), "white", outline=MID_BLUE, width=3, radius=14)
    centered(draw, (515, 535, 945, 675), "模式定位逻辑 + 检查验证逻辑\n检查器结构约束", ft_small)
    arrow(draw, (890, 455), (770, 535), MID_BLUE, 5, 15)

    # Stage 2 details.
    rounded_box(draw, (1160, 290, 1675, 470), "white", outline=MID_BLUE, width=3, radius=14)
    centered(draw, (1160, 290, 1675, 470), "检测子逻辑语义检索\n模式匹配API / AST操作API\n匹配元操作 / 检查验证元操作", ft_small)
    rounded_box(draw, (1190, 555, 1645, 675), "white", outline=MID_BLUE, width=3, radius=14)
    centered(draw, (1190, 555, 1645, 675), "AST相关性重排\n输出API上下文与实现片段", ft_small)
    arrow(draw, (1417, 470), (1417, 555), MID_BLUE, 5, 15)

    # Stage 3 details.
    rounded_box(draw, (1845, 285, 2075, 420), "white", outline=GREEN, width=3, radius=14)
    centered(draw, (1845, 285, 2075, 420), "完整候选\n检查器生成", ft_small)
    rounded_box(draw, (2120, 285, 2325, 420), "white", outline=ORANGE, width=3, radius=14)
    centered(draw, (2120, 285, 2325, 420), "编译验证", ft_body)
    rounded_box(draw, (2365, 285, 2500, 420), "white", outline=ORANGE, width=3, radius=14)
    centered(draw, (2365, 285, 2500, 420), "测试\n验证", ft_body)
    arrow(draw, (2075, 352), (2120, 352), GREEN, 5, 15)
    arrow(draw, (2325, 352), (2365, 352), GREEN, 5, 15)
    rounded_box(draw, (2050, 555, 2340, 665), "white", outline=GREEN, width=3, radius=14)
    centered(draw, (2050, 555, 2340, 665), "最终检查器", ft_body, GREEN, stroke=1)
    arrow(draw, (2432, 420), (2340, 580), GREEN, 5, 15)

    # Main arrows.
    arrow(draw, (330, 450), (420, 450), BLUE, 7, 22)
    arrow(draw, (1040, 450), (1110, 450), BLUE, 7, 22)
    arrow(draw, (1725, 450), (1795, 450), BLUE, 7, 22)

    # Feedback labels and paths.
    draw.text((2000, 455), "编译失败：修正代码及补充接口知识", font=ft_note, fill=RED)
    poly_arrow(draw, [(2220, 445), (2220, 495), (1960, 495), (1960, 420)], RED, 4, 15)
    draw.text((1400, 785), "漏报 / 误报：调整检测逻辑并重新检索", font=ft_note, fill=PURPLE)
    poly_arrow(draw, [(2410, 445), (2410, 820), (735, 820), (735, 745)], PURPLE, 5, 18)

    # Offline knowledge foundation.
    rounded_box(draw, (625, 920, 1965, 1110), LIGHT_GRAY, outline="#888888", width=3, radius=20)
    centered(draw, (650, 935, 920, 1090), "离线知识基础", ft_stage, GRAY, stroke=1)
    rounded_box(draw, (940, 942, 1260, 1008), "white", outline="#888888", width=3, radius=14)
    centered(draw, (940, 942, 1260, 1008), "目标框架API", ft_small)
    rounded_box(draw, (940, 1025, 1260, 1092), "white", outline="#888888", width=3, radius=14)
    centered(draw, (940, 1025, 1260, 1092), "已有检查器源码", ft_small)
    rounded_box(draw, (1530, 955, 1905, 1075), "white", outline="#888888", width=3, radius=14)
    centered(draw, (1530, 955, 1905, 1075), "API库与元操作库", ft_small)
    arrow(draw, (1260, 975), (1530, 1012), GRAY, 4, 14)
    arrow(draw, (1260, 1058), (1530, 1018), GRAY, 4, 14)
    # Dashed link from foundation to stage 2.
    for y in range(920, 745, -22):
        draw.line([(1650, y), (1650, max(y-12, 745))], fill=GRAY, width=3)
    draw.polygon([(1650, 745), (1642, 762), (1658, 762)], fill=GRAY)

    img.save(FIGURE, quality=95)


def remove_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def format_body(paragraph):
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def build_document():
    make_figure()
    doc = Document(SOURCE)

    # Capture the original nodes before inserting/replacing content.
    chapter_heading = doc.paragraphs[1]
    old_text = doc.paragraphs[2]
    old_blank = doc.paragraphs[3]
    old_figure = doc.paragraphs[4]
    old_caption = doc.paragraphs[5]
    chapter2_heading = doc.paragraphs[6]

    chapter_heading.text = "1 技术概述"
    chapter_heading.style = doc.styles["Normal"]
    chapter_heading.paragraph_format.space_before = Pt(12)
    chapter_heading.paragraph_format.space_after = Pt(8)
    for run in chapter_heading.runs:
        run.font.name = "黑体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(31, 78, 121)

    paragraphs = [
        "本项目面向定制静态代码检查器开发过程中规则转换困难、分析框架接口复杂以及调试周期较长等问题，研究测试用例驱动的静态代码检查器生成技术。该技术以自然语言规则描述和带有预期结果的测试用例为输入，自动生成能够接入目标静态分析框架的检查器代码。当前工作主要面向C/C++代码，并以clang-tidy作为检查器的运行和验证载体。",
        "如图1所示，检查器生成过程由三个阶段组成。第一阶段以违规用例为种子，结合规则描述和测试代码的抽象语法树提取待检查的代码模式，并将检查过程拆分为模式定位逻辑和检查验证逻辑，在此基础上形成目标节点、节点关系、上下文条件、排除条件和诊断位置等检查器结构约束。第二阶段以各项检测子逻辑为查询，从模式匹配API、AST节点操作API、匹配元操作和检查验证元操作等知识库中检索相关接口及实现片段，并结合测试用例中实际出现的AST节点类型对检索结果进行筛选，形成与当前代码模式相匹配的API上下文。第三阶段综合规则描述、测试用例、抽象语法树、检测逻辑、结构约束和检索结果，在目标检查器模板基础上生成完整的候选检查器。",
        "候选检查器生成后依次进行编译验证和测试验证。编译失败时，根据错误信息修正接口调用和代码结构，并补充检索所需的API知识；编译通过后，使用完整测试套件检查候选结果。违规用例未产生告警时，利用该用例补充检测模式；合规用例产生告警时，利用该用例增加排除条件。每次修正后重新执行测试，使新增逻辑不破坏已经通过的用例。",
        "截至中期检查阶段，项目已完成上述三阶段生成流程以及由编译反馈和测试反馈构成的迭代闭环，完成了模式匹配API、AST节点操作API、匹配元操作和检查验证元操作等知识的收集与组织。现有阶段性数据包含10条C/C++规则和200个测试用例，生成检查器的平均测试通过率为99.5%，违规用例召回率为99.0%。后续将继续扩大规则验证范围，并重点完善复杂规则的生成稳定性和处理效率。",
    ]

    # Insert new text and figure immediately before the old figure position.
    for text in paragraphs[:3]:
        p = old_figure.insert_paragraph_before(text)
        format_body(p)

    old_figure.clear()
    old_figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    old_figure.paragraph_format.space_before = Pt(5)
    old_figure.paragraph_format.space_after = Pt(3)
    old_figure.add_run().add_picture(str(FIGURE), width=Inches(6.05))

    old_caption.text = "图1  AutoChecker测试用例驱动的静态代码检查器生成技术路线"
    old_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    old_caption.paragraph_format.space_before = Pt(0)
    old_caption.paragraph_format.space_after = Pt(8)
    for run in old_caption.runs:
        run.font.name = "宋体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(9)

    p = chapter2_heading.insert_paragraph_before(paragraphs[3])
    format_body(p)

    remove_paragraph(old_text)
    remove_paragraph(old_blank)

    doc.core_properties.title = "测试用例驱动的静态代码检查器生成技术"
    doc.core_properties.subject = "项目中期检查材料（第一章修改版）"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
