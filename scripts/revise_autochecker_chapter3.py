#!/usr/bin/env python3
"""Rewrite Chapter 3 in-place as a method-focused midterm-report chapter."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "docs/测试用例驱动的静态代码检查器生成技术-第一二章修改版.docx"
ASSET_DIR = ROOT / "docs/autochecker_chapter3_assets"
FONT_PATH = Path("/tmp/NotoSansCJKsc-Regular.otf")
FIGURE_KB = ASSET_DIR / "图2_静态分析知识库构建方法.png"
FIGURE_RETRIEVAL = ASSET_DIR / "图3_检测逻辑驱动的分层检索方法.png"

BLUE = "#1F4E79"
MID_BLUE = "#5B9BD5"
LIGHT_BLUE = "#D9EAF7"
PALE_BLUE = "#EDF4FA"
GREEN = "#70AD47"
LIGHT_GREEN = "#E2F0D9"
ORANGE = "#ED7D31"
LIGHT_ORANGE = "#FCE4D6"
PURPLE = "#7030A0"
GRAY = "#666666"
LIGHT_GRAY = "#F2F2F2"
WHITE = "#FFFFFF"


def image_font(size):
    if not FONT_PATH.exists():
        raise FileNotFoundError(f"缺少中文字体：{FONT_PATH}")
    return ImageFont.truetype(str(FONT_PATH), size=size)


def rounded(draw, box, fill, outline=BLUE, width=4, radius=20):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, ft, fill="#222222", spacing=7, stroke=0):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=ft, spacing=spacing,
                                   align="center", stroke_width=stroke)
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(((x1+x2-width)/2, (y1+y2-height)/2), text, font=ft,
                        fill=fill, spacing=spacing, align="center", stroke_width=stroke)


def arrow(draw, start, end, fill=BLUE, width=6, head=18):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2-x1) >= abs(y2-y1):
        sign = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2-sign*head, y2-head//2), (x2-sign*head, y2+head//2)]
    else:
        sign = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2-head//2, y2-sign*head), (x2+head//2, y2-sign*head)]
    draw.polygon(points, fill=fill)


def make_kb_figure():
    ASSET_DIR.mkdir(exist_ok=True)
    image = Image.new("RGB", (2500, 1120), "white")
    draw = ImageDraw.Draw(image)
    # Figures are inserted into a portrait A4 page at about 15 cm wide.  Use
    # deliberately large raster text so labels remain readable after scaling.
    title = image_font(62)
    heading = image_font(50)
    body = image_font(44)
    small = image_font(38)

    center_text(draw, (0, 15, 2500, 90), "静态分析知识的收集、抽象与知识库构建",
                title, BLUE, stroke=1)

    # Sources.
    rounded(draw, (60, 180, 500, 900), LIGHT_GRAY, outline="#888888")
    center_text(draw, (75, 205, 485, 285), "知识来源", heading, GRAY, stroke=1)
    rounded(draw, (115, 340, 445, 510), WHITE, outline="#888888", width=3)
    center_text(draw, (115, 340, 445, 510), "目标静态分析框架\n接口声明与说明", small)
    rounded(draw, (115, 610, 445, 780), WHITE, outline="#888888", width=3)
    center_text(draw, (115, 610, 445, 780), "已有高质量\n检查器实现", body)

    # Abstraction.
    rounded(draw, (650, 180, 1250, 900), LIGHT_BLUE)
    center_text(draw, (670, 205, 1230, 285), "知识抽取与语义化", heading, BLUE, stroke=1)
    rounded(draw, (710, 335, 1190, 525), WHITE, outline=MID_BLUE, width=3)
    center_text(draw, (710, 335, 1190, 525), "提取接口名称、签名、参数\n与返回类型，形成面向功能的\n语义描述", small, spacing=5)
    rounded(draw, (710, 600, 1190, 790), WHITE, outline=MID_BLUE, width=3)
    center_text(draw, (710, 600, 1190, 790), "分解匹配过程与检查过程\n提取可复用Meta OP\n及实现片段", small, spacing=5)
    arrow(draw, (500, 425), (710, 425), BLUE)
    arrow(draw, (500, 695), (710, 695), BLUE)

    # Knowledge forms.
    rounded(draw, (1400, 180, 1960, 900), LIGHT_GREEN, outline=GREEN)
    center_text(draw, (1420, 205, 1940, 285), "分层知识表示", heading, GREEN, stroke=1)
    rounded(draw, (1450, 325, 1910, 520), WHITE, outline=GREEN, width=3)
    center_text(draw, (1450, 325, 1910, 520), "API数据库\n模式匹配API子集\nAST节点操作API子集", small)
    rounded(draw, (1450, 595, 1910, 790), WHITE, outline=GREEN, width=3)
    center_text(draw, (1450, 595, 1910, 790), "Meta OP数据库\n匹配Meta OP子集\n检查验证Meta OP子集", small)
    arrow(draw, (1250, 425), (1450, 425), GREEN)
    arrow(draw, (1250, 695), (1450, 695), GREEN)

    # Database construction.
    rounded(draw, (2100, 180, 2440, 900), LIGHT_ORANGE, outline=ORANGE)
    center_text(draw, (2115, 205, 2425, 285), "知识库构建", heading, ORANGE, stroke=1)
    rounded(draw, (2145, 350, 2395, 500), WHITE, outline=ORANGE, width=3)
    center_text(draw, (2145, 350, 2395, 500), "规范化与\n融合去重", body)
    rounded(draw, (2145, 590, 2395, 740), WHITE, outline=ORANGE, width=3)
    center_text(draw, (2145, 590, 2395, 740), "语义索引与\n来源关联", body)
    arrow(draw, (1960, 430), (2145, 430), ORANGE)
    arrow(draw, (1960, 690), (2145, 665), ORANGE)

    center_text(draw, (400, 970, 2110, 1075),
                "知识条目以“功能语义—框架载体”为核心组织，为在线检索提供统一入口",
                small, BLUE, stroke=1)
    image.save(FIGURE_KB, quality=95)


def make_retrieval_figure():
    image = Image.new("RGB", (2500, 1180), "white")
    draw = ImageDraw.Draw(image)
    title = image_font(62)
    heading = image_font(50)
    body = image_font(42)
    small = image_font(38)

    center_text(draw, (0, 15, 2500, 90), "检测逻辑驱动的分层知识检索方法",
                title, BLUE, stroke=1)

    # Query logic.
    rounded(draw, (60, 180, 485, 890), LIGHT_GRAY, outline="#888888")
    center_text(draw, (75, 205, 470, 285), "检测逻辑", heading, GRAY, stroke=1)
    rounded(draw, (110, 345, 435, 510), WHITE, outline=MID_BLUE, width=3)
    center_text(draw, (110, 345, 435, 510), "模式定位逻辑\n目标节点与\n结构关系", small, spacing=3)
    rounded(draw, (110, 630, 435, 795), WHITE, outline=GREEN, width=3)
    center_text(draw, (110, 630, 435, 795), "检查验证逻辑\n属性、上下文\n与诊断", small, spacing=3)

    # Routed stores.
    rounded(draw, (650, 180, 1240, 890), PALE_BLUE)
    center_text(draw, (670, 205, 1220, 285), "按逻辑类型定向检索", heading, BLUE, stroke=1)
    rounded(draw, (705, 305, 1185, 555), WHITE, outline=MID_BLUE, width=3)
    center_text(draw, (705, 305, 1185, 555), "API数据库\n模式匹配API子集\nMeta OP数据库\n匹配Meta OP子集", small, spacing=2)
    rounded(draw, (705, 570, 1185, 840), WHITE, outline=GREEN, width=3)
    center_text(draw, (705, 570, 1185, 840), "API数据库\nAST节点操作API子集\nMeta OP数据库\n检查验证Meta OP子集", small, spacing=2)
    arrow(draw, (485, 425), (705, 425), MID_BLUE)
    arrow(draw, (485, 712), (705, 712), GREEN)

    # Retrieval and reranking.
    rounded(draw, (1410, 180, 1925, 890), LIGHT_BLUE)
    center_text(draw, (1430, 205, 1905, 285), "候选召回与重排", heading, BLUE, stroke=1)
    rounded(draw, (1460, 330, 1875, 500), WHITE, outline=MID_BLUE, width=3)
    center_text(draw, (1460, 330, 1875, 500), "计算检测逻辑与\n知识语义的相关性", body)
    rounded(draw, (1460, 580, 1875, 800), WHITE, outline=PURPLE, width=3)
    center_text(draw, (1460, 580, 1875, 800), "结合测试用例AST\n进行结构相关性\n重排", body, spacing=4)
    arrow(draw, (1240, 425), (1460, 415), BLUE)
    arrow(draw, (1240, 712), (1460, 690), BLUE)
    arrow(draw, (1667, 500), (1667, 580), PURPLE)

    # Output context.
    rounded(draw, (2070, 180, 2440, 890), LIGHT_GREEN, outline=GREEN)
    center_text(draw, (2085, 205, 2425, 285), "检索输出", heading, GREEN, stroke=1)
    rounded(draw, (2120, 335, 2390, 500), WHITE, outline=GREEN, width=3)
    center_text(draw, (2120, 335, 2390, 500), "接口说明\n与方法签名", body)
    rounded(draw, (2120, 610, 2390, 775), WHITE, outline=GREEN, width=3)
    center_text(draw, (2120, 610, 2390, 775), "可复用的\n实现片段", body)
    arrow(draw, (1925, 425), (2120, 420), GREEN)
    arrow(draw, (1925, 690), (2120, 690), GREEN)

    rounded(draw, (630, 980, 1930, 1110), LIGHT_ORANGE, outline=ORANGE, width=3)
    center_text(draw, (650, 990, 1910, 1100),
                "编译或测试反馈产生新的知识需求时\n重新进入相应知识库进行补充检索",
                small, ORANGE, spacing=3, stroke=1)
    # The feedback is an external input to retrieval, not a result produced by
    # the retrieval output itself.
    arrow(draw, (630, 1045), (280, 890), ORANGE, 5, 17)
    image.save(FIGURE_RETRIEVAL, quality=95)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_run_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_numbers(doc):
    """Add a centered dynamic page-number field to every section footer."""
    for section in doc.sections:
        section.different_first_page_header_footer = False
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        prefix = paragraph.add_run("第 ")
        set_run_font(prefix, size=9, color="666666")

        field_run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = "PAGE"
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        result = OxmlElement("w:t")
        result.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        field_run._r.extend([begin, instruction, separate, result, end])
        set_run_font(field_run, size=9, color="666666")

        suffix = paragraph.add_run(" 页")
        set_run_font(suffix, size=9, color="666666")

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def format_heading(paragraph, level):
    sizes = {1: 16, 2: 13.5, 3: 11.5}
    before = {1: 14, 2: 12, 3: 9}
    after = {1: 8, 2: 6, 3: 4}
    paragraph.paragraph_format.space_before = Pt(before[level])
    paragraph.paragraph_format.space_after = Pt(after[level])
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(run, "黑体", sizes[level], True,
                     "1F4E79" if level < 3 else "365F91")


def add_heading(doc, text, level):
    p = doc.add_paragraph(text)
    p.style = "Normal"
    format_heading(p, level)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run)
    return p


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(str(text))
    set_run_font(run, size=8.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "A6A6A6")
        borders.append(node)
    tblpr.append(borders)
    trpr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], "1F4E79")
        set_cell(table.rows[0].cells[i], header, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if ridx % 2:
                shade_cell(cells[i], "EDF4FA")
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table_caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run, size=9, color="666666")
    return p


def add_figure(doc, path, label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(6.0))
    cp = doc.add_paragraph(label)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    for run in cp.runs:
        set_run_font(run, size=9, color="666666")


def write_chapter(doc):
    add_body(doc, "本章给出测试用例驱动的静态代码检查器生成方法。方法以规则描述和带有预期结果的测试用例为输入，通过代码模式分析形成检测逻辑，以预先构建的静态分析知识库补充框架知识，在此基础上生成候选检查器，并利用编译结果和测试结果进行迭代修正。整个过程对应图1中的三个阶段：检测逻辑建模、知识检索以及候选检查器生成与反馈验证。其中，知识库建设作为离线基础，为第二阶段的在线检索提供支持。")
    add_body(doc, "该方法的基本思路是把一次完整的检查器生成任务拆成若干可验证的中间问题。规则和测试用例首先被转换为与程序结构相关的检测步骤；检测步骤随后被转换为知识查询，用于定位可用的分析接口及实现方式；候选检查器生成后，则由编译器和测试套件提供确定性反馈。这样可以减少从自然语言直接生成完整检查器时产生的语义偏差和框架接口错误。")

    add_heading(doc, "3.1 阶段1：基于代码模式的检测逻辑建模", 2)
    add_heading(doc, "3.1.1 规则与测试用例的联合约束", 3)
    add_body(doc, "规则描述与测试用例分别承担语义约束和行为约束。规则描述给出检查目标、适用范围以及允许或禁止的代码行为，但通常不会完整列出所有语法形式；测试用例把规则落实为具体代码，并通过预期告警或预期静默给出可执行的判断标准。方法将预期产生告警的代码称为违规用例，将预期不产生告警的代码称为合规用例。")
    add_body(doc, "初始生成以一个具有代表性的违规用例作为种子。该用例提供最基本的违例结构，使系统可以先建立一条能够工作的检测路径。候选检查器形成后，再使用完整测试套件检验其覆盖范围：其他违规用例用于发现尚未覆盖的代码模式，合规用例用于识别需要排除的合法情形。因而，测试用例既参与初始建模，也参与后续边界修正。")

    add_heading(doc, "3.1.2 基于抽象语法树的代码模式分析", 3)
    add_body(doc, "静态代码检查通常面向程序的语法和语义结构，而不是源代码的表面字符串。为此，方法首先对种子用例进行抽象语法树分析，识别与规则相关的声明、语句和表达式节点，并确定目标节点与其父节点、子节点或上下文节点之间的关系。抽象语法树还提供节点类型、源代码位置和隐式结构等信息，可用于区分表面形式相近但语义不同的代码。")
    add_body(doc, "代码模式分析的目标不是复现整个语法树，而是提取足以支撑规则判断的局部结构。对于简单规则，局部结构可能只包含一个目标节点及其属性；对于上下文相关规则，还需要描述目标节点所在的控制结构、作用域、声明关系或调用关系。通过这种局部化表示，规则语义被转化为后续匹配和验证可以直接使用的程序结构约束。")

    add_heading(doc, "3.1.3 检测逻辑集合与检查器结构约束", 3)
    add_body(doc, "在代码模式分析的基础上，方法将检查过程拆分为模式定位逻辑和检查验证逻辑。模式定位逻辑用于在程序中筛选候选节点，主要描述节点类型、节点之间的结构关系以及需要保留的节点绑定；检查验证逻辑用于对候选节点进行进一步判断，主要描述节点属性读取、上下文条件、排除条件和诊断位置。两类逻辑形成有序的检测逻辑集合。")
    add_body(doc, "检测逻辑集合进一步形成图1所示的检查器结构约束。该约束包括目标节点、节点绑定、结构条件、上下文条件、排除条件和诊断位置等要素。它并非对测试用例的直接复制，而是对规则成立条件的抽象，并作为独立的中间表示约束后续知识检索和候选检查器生成。")
    add_body(doc, "当后续测试发现漏报时，新的违规模式被补充到模式定位逻辑或检查验证逻辑中；当测试发现误报时，合规用例所体现的差异被转化为新的排除条件。由此，检测逻辑不是一次生成后固定不变，而是在测试反馈约束下逐步完善。")

    add_heading(doc, "3.2 离线支撑：API数据库与Meta OP数据库的收集和构建", 2)
    add_heading(doc, "3.2.1 知识来源与构建目标", 3)
    add_body(doc, "静态分析框架中的知识可以分为两个层面。第一个层面是框架明确提供的程序分析接口，它回答某个接口可以操作哪类节点、接收哪些参数以及返回什么结果；第二个层面是已有检查器中沉淀的接口使用经验，它回答多个接口如何组合才能完成一项具体检测动作。若只收集第一类知识，生成模型仍需自行推断接口组合；若只收集第二类知识，又容易照搬既有规则中的局部实现。基于这一认识，本项目将知识库统一划分为API数据库和Meta OP数据库。")
    add_body(doc, "API数据库以目标静态分析框架的接口声明和接口说明为主要来源，保存功能语义与接口签名之间的对应关系。Meta OP数据库以已有高质量检查器实现为主要来源，保存检测操作语义与实现片段之间的对应关系。两类数据库不是相互替代的关系：API数据库提供较细粒度、类型明确的基础能力，Meta OP数据库提供较高层次、可以直接参考的组合经验。在线检索同时使用二者，以兼顾接口准确性和实现完整性。")
    add_body(doc, "知识库建设的目标并非保存完整框架文档，也不是复制已有检查器，而是建立从规则检测意图到框架实现方式的中间映射。为此，两类数据库中的知识条目均采用“功能语义—框架载体”的组织方式。功能语义用于与检测逻辑进行比较；框架载体在API数据库中表现为接口名称和签名，在Meta OP数据库中表现为一段能够完成相应操作的代码。")

    add_heading(doc, "3.2.2 API数据库的收集与构建", 3)
    add_body(doc, "API数据库的收集对象包括两类接口。一类是用于描述程序结构的模式匹配接口，包括节点匹配、属性收窄和结构遍历等操作；另一类是抽象语法树节点提供的操作接口，包括节点属性读取、声明关系获取、类型判断、源码位置访问和上下文导航等操作。前一类构成模式匹配API子集，主要服务于候选节点定位；后一类构成AST节点操作API子集，主要服务于命中后的检查验证。")
    add_body(doc, "模式匹配API的收集以框架提供的匹配器定义和说明为依据。对于每个接口，首先记录接口名称、适用节点类型、输入参数、返回的匹配器类型以及功能说明，然后将原始说明归纳为面向检测任务的语义描述。例如，接口binaryOperator用于匹配二元运算表达式，hasAncestor用于约束目标节点具有某类祖先节点，isAssignmentOperator用于判断运算节点是否属于赋值运算。经过语义化后，它们可以分别与“定位二元运算”“限制节点所在上下文”和“筛选赋值运算”等检测逻辑建立联系。")
    add_body(doc, "AST节点操作API的收集以框架中的节点类型和成员方法为依据。收集过程以节点类型为边界，提取方法名称、参数、返回类型和完整签名，并根据方法作用生成统一的功能描述。例如，BinaryOperator::getOperatorLoc()可描述为“获取二元运算符的源码位置”，NamedDecl::getNameAsString()可描述为“获取具名声明的名称”，CallExpr::getDirectCallee()可描述为“获取调用表达式直接对应的函数声明”。这些描述比原始方法名更接近规则和检测逻辑使用的语言。")
    add_body(doc, "接口语义化是API数据库构建中的关键步骤。静态分析接口名称通常包含节点类型、缩写和框架内部术语，规则描述则使用“判断”“获取”“查找”“位于”等自然语言动作。方法将接口转换为动作—对象形式的功能描述，并把描述与原始签名绑定。这样，检索阶段使用功能描述完成语义匹配，候选生成阶段使用原始签名约束参数类型、返回结果和调用形式。")
    add_body(doc, "表1给出了API数据库条目的示例。示例中的功能描述是面向检索的语义表示，接口及签名是提供给候选生成阶段的框架知识。一个接口可能适用于多个节点类型，或者存在不同参数形式，此时分别保留其类型约束和方法形式，避免只根据接口名称进行不加区分的复用。")
    add_table_caption(doc, "表1  API数据库典型知识条目")
    add_table(doc, ["知识子集", "功能语义示例", "接口或签名示例", "在检查器中的作用"], [
        ("模式匹配API", "匹配二元运算表达式", "binaryOperator(...) → Matcher<Stmt>", "建立候选表达式集合"),
        ("模式匹配API", "匹配具有指定祖先的节点", "hasAncestor(Matcher<...>)", "表达目标节点的上下文关系"),
        ("模式匹配API", "筛选赋值运算", "isAssignmentOperator()", "对候选运算节点进行属性收窄"),
        ("AST节点操作API", "获取二元运算符的源码位置", "BinaryOperator::getOperatorLoc()", "确定告警位置"),
        ("AST节点操作API", "获取调用表达式对应的函数声明", "CallExpr::getDirectCallee()", "取得被调用函数并继续判断"),
        ("AST节点操作API", "获取具名声明的名称", "NamedDecl::getNameAsString()", "读取声明名称"),
    ], [2.7, 4.2, 4.5, 3.8])
    add_body(doc, "API数据库中的知识粒度以单个接口或单个接口族为主。该粒度有利于准确控制节点类型和参数，但无法完整表示一段检查逻辑所需的接口组合。因此，API数据库主要解决“有哪些可用接口”和“接口如何正确调用”的问题，更高层次的实现组合由Meta OP数据库补充。")

    add_heading(doc, "3.2.3 Meta OP数据库的收集与构建", 3)
    add_body(doc, "Meta OP是对已有检查器中重复出现的程序分析操作的抽象。这里的“Meta”表示该操作已经脱离某一条具体规则，但仍保留完成检测任务所需的框架实现。例如，“匹配所有函数定义”“获取已绑定的函数声明并检查其有效性”和“在目标节点位置输出诊断”都可以作为Meta OP。它们比单个API具有更完整的语义，又比完整检查器具有更好的复用性。")
    add_body(doc, "Meta OP数据库的原始语料来自已有高质量检查器源码。这些检查器已经经过框架工程的编译和测试，其中包含大量未直接写入API文档的使用经验，例如多个匹配器的嵌套方式、节点绑定和取回之间的对应关系、空节点防护以及诊断位置选择。收集过程首先识别检查器的主要处理单元，再将处理单元划分为模式匹配过程和检查验证过程。")
    add_body(doc, "模式匹配过程描述候选节点如何被定位和绑定。方法从该部分识别节点匹配、属性收窄、上下文约束、排除条件和组合匹配等操作，并将其归纳为匹配Meta OP。例如，从functionDecl(isDefinition())可以抽象出“匹配属于定义的函数声明”；从functionDecl(isDefinition(), unless(isDeleted()))可以抽象出“匹配已经定义且未被删除的函数声明”。抽象描述不包含原检查器的规则名称，因此可以被其他需要定位函数定义的规则检索和复用。")
    add_body(doc, "检查验证过程描述候选节点命中后如何完成进一步判断。方法从该部分识别绑定节点获取、有效性检查、属性读取、上下文遍历、提前返回、诊断输出和修复提示等操作，并将其归纳为检查验证Meta OP。例如，一段代码可能先从匹配结果中取得FunctionDecl节点，在节点不存在、没有显式原型或属于模板实例时返回。该片段可以被抽象为“获取已匹配的函数声明并排除无效或不适用节点”。")
    add_body(doc, "Meta OP抽取时需要控制操作粒度。粒度过小会退化为单个API，无法提供组合经验；粒度过大会保留原规则的大量特定逻辑，难以在其他规则中复用。方法以“能够用一句独立的检测动作描述，并能够由一段相对完整的代码实现”为基本划分原则。对于同时包含多个语义目标的片段，需要继续拆分；对于只有语法作用而缺少独立检测意义的片段，则与其上下文合并。")
    add_body(doc, "每条Meta OP记录包括功能语义、实现片段、操作类型和知识来源。功能语义使用与检测逻辑相近的表达，代码片段保留关键接口组合和数据依赖；原规则名称、局部变量名称和具体诊断文本等非必要信息在语义层面被弱化。随后对不同检查器中抽取的Meta OP进行融合和去重，使功能相同但写法略有差异的操作形成统一的检索入口，同时保留具有不同节点约束或上下文条件的实现变体。")
    add_body(doc, "表2给出了Meta OP数据库中的典型条目。与表1中的API数据库条目相比，Meta OP数据库条目通常已经包含多个接口之间的组合、节点绑定或有效性判断，因此能够为候选检查器提供更接近最终实现的参考。")
    add_table_caption(doc, "表2  Meta OP数据库典型知识条目")
    add_table(doc, ["Meta OP子集", "操作语义示例", "实现片段示例", "可复用信息"], [
        ("匹配Meta OP", "匹配属于定义的函数声明", "functionDecl(isDefinition())", "节点匹配与属性收窄的组合"),
        ("匹配Meta OP", "匹配已定义且未删除的函数声明", "functionDecl(isDefinition(), unless(isDeleted()))", "正向条件与排除条件的组合"),
        ("检查验证Meta OP", "取得已绑定的函数声明，不存在时结束检查", "getNodeAs<FunctionDecl>(\"function\"); if (!Function) return;", "绑定名称、节点类型与空值防护"),
        ("检查验证Meta OP", "取得已绑定的字符串字面量并验证有效性", "getNodeAs<StringLiteral>(\"strlit\"); if (!SL) return;", "节点获取和有效性检查模式"),
    ], [2.8, 4.3, 5.4, 3.0])
    add_body(doc, "API数据库和Meta OP数据库的差别可以概括为：API数据库描述框架提供的基础能力，Meta OP数据库描述这些能力在检查器中的典型使用方式。例如，对于“匹配函数定义”这一检测逻辑，API数据库可能返回functionDecl和isDefinition两个接口的说明，Meta OP数据库则可以返回functionDecl(isDefinition())这一组合片段。两类结果共同使用时，既能够核对接口类型，也能够参考已经成立的组合形式。")

    add_heading(doc, "3.2.4 两类数据库的组织、去重与索引", 3)
    add_body(doc, "经过收集和抽取后，两类知识分别进入API数据库和Meta OP数据库。API数据库内部进一步区分模式匹配API子集与AST节点操作API子集；Meta OP数据库内部进一步区分匹配Meta OP子集与检查验证Meta OP子集。该组织与阶段1中的模式定位逻辑和检查验证逻辑相对应，为后续定向检索提供明确边界。")
    add_table_caption(doc, "表3  API数据库与Meta OP数据库的构成及作用")
    add_table(doc, ["数据库", "知识子集", "主要来源", "知识形式", "主要作用"], [
        ("API数据库", "模式匹配API", "目标静态分析框架的匹配接口", "功能描述—接口签名", "支持目标节点和结构关系表达"),
        ("API数据库", "AST节点操作API", "目标静态分析框架的AST类型与方法", "功能描述—方法签名", "支持节点属性与上下文获取"),
        ("Meta OP数据库", "匹配Meta OP", "已有高质量检查器的匹配过程", "操作描述—匹配实现片段", "提供匹配接口的组合方式"),
        ("Meta OP数据库", "检查验证Meta OP", "已有高质量检查器的验证过程", "操作描述—验证实现片段", "提供节点获取、判断和诊断方式"),
    ], [2.5, 3.1, 3.6, 3.6, 3.5])
    add_body(doc, "数据库构建首先统一功能描述的表达方式，使不同来源中含义相同的知识能够被共同检索。随后根据功能语义、接口签名或代码结构识别重复项。对于完全等价的知识进行合并，对于语义相同但适用节点类型、输入参数或上下文不同的条目保留为不同变体，避免去重过程丢失有效约束。每条知识同时保留来源关联，以便在出现歧义时回溯原始接口或检查器上下文。")
    add_body(doc, "完成规范化和融合后，两类数据库分别对知识条目的功能语义建立检索索引。索引只作用于语义描述，接口签名和实现片段作为命中后的返回内容。这种“语义检索、框架载体返回”的方式将自然语言规则与框架代码分开处理：检索过程关注功能是否相近，生成过程关注接口和片段如何正确使用。")
    add_body(doc, "截至中期检查阶段，API数据库收录670个模式匹配API；其中，AST节点操作API子集覆盖1836类AST节点，共包含15574个成员方法。Meta OP数据库收录1822条匹配Meta OP和1668条检查验证Meta OP。在线生成时不会把全部数据库内容直接提供给生成模型，而是根据当前检测逻辑从两个数据库中按需选择少量相关知识。")
    add_figure(doc, FIGURE_KB, "图2  静态分析知识的收集与知识库构建方法")

    add_heading(doc, "3.3 阶段2：基于检测逻辑的知识检索", 2)
    add_heading(doc, "3.3.1 检索查询构造与知识路由", 3)
    add_body(doc, "在线检索以阶段1得到的检测逻辑为基本查询单元，而不是直接使用完整规则描述。完整规则通常同时包含目标节点、上下文关系、排除条件和诊断要求，若以整段文本进行一次检索，不同意图容易相互干扰。将规则分解后的每条检测逻辑分别查询，可以使检索结果围绕一个明确的程序分析动作展开。")
    add_body(doc, "知识路由依据检测逻辑在检查器中的作用完成。模式定位逻辑被发送到API数据库的模式匹配API子集和Meta OP数据库的匹配Meta OP子集，用于获取节点匹配方式及组合片段；检查验证逻辑被发送到API数据库的AST节点操作API子集和Meta OP数据库的检查验证Meta OP子集，用于获取节点属性访问、上下文判断和诊断实现。该分层方式使功能语义相近但使用阶段不同的知识保持区分。")
    add_body(doc, "以“禁止在条件表达式中直接使用赋值运算”为例，模式定位逻辑可以分解为“匹配赋值运算节点”和“约束该节点位于条件语句中”，这两条逻辑主要查询模式匹配API和匹配Meta OP；检查验证逻辑可以分解为“获取运算符源码位置”和“在该位置输出诊断”，这两条逻辑主要查询AST节点操作API和检查验证Meta OP。分解后的查询保持原规则中的依赖关系，但每次检索只处理一个明确目标。")
    add_body(doc, "对于同时涉及模式定位和语义判断的逻辑，方法不会简单地将其固定到单一知识子集，而是先根据主要目的完成初始路由，再在检索结果不足时扩展到另一类知识。这样既维持分层检索的针对性，也避免由于逻辑表述边界不清而遗漏必要知识。")

    add_heading(doc, "3.3.2 基于语义相关性的候选召回", 3)
    add_body(doc, "为消除规则描述与框架接口之间的词汇差异，方法采用语义表示比较检测逻辑和知识描述。设第i条检测逻辑为lᵢ，知识库中的第j条知识为kⱼ，二者的语义相关性由向量表示之间的余弦相似度衡量，即s_sem(lᵢ,kⱼ)=cos(E(lᵢ),E(kⱼ))。每条检测逻辑分别从对应知识库中召回一组高相关候选。")
    add_body(doc, "语义召回可以识别措辞不同但功能相近的知识。例如，“取得调用表达式对应的函数声明”和“访问被调用函数”在文本上并不完全一致，但指向相近的程序分析操作。通过功能语义匹配，检索不再依赖接口名称的精确出现，也减少了生成模型凭记忆猜测框架接口的需要。")
    add_body(doc, "API数据库和Meta OP数据库采用相同的语义入口，但返回内容不同。查询API数据库时，候选结果包含接口说明、方法签名和适用节点类型；查询Meta OP数据库时，候选结果包含操作描述和实现片段。例如，查询“匹配函数定义”时，API数据库可能分别返回functionDecl和isDefinition的知识，Meta OP数据库则可能返回functionDecl(isDefinition())的组合实现。两类候选在后续上下文组织阶段进行互补。")
    add_body(doc, "对于由多条步骤组成的检测逻辑集合，系统分别执行召回，再合并各步骤的候选结果。合并时保留知识与原检测步骤之间的对应关系，使生成阶段能够判断某项接口或片段是用于目标节点定位、属性检查还是诊断输出，而不是将所有候选无差别堆叠。")

    add_heading(doc, "3.3.3 测试用例AST驱动的结构相关性重排", 3)
    add_body(doc, "语义相关性只表明知识在功能描述上接近，不能保证它适用于当前测试代码。静态分析框架中常有多个名称或用途相近的接口，但它们面向不同节点类型。为提高候选知识与样例结构的一致性，方法进一步利用测试用例AST中的节点类型对初始候选进行重排。")
    add_body(doc, "重排过程优先保留显式涉及当前样例节点类型的知识，同时保留能够描述父子、祖先和后代关系的通用遍历操作。对于未包含明显节点类型但语义高度相关的结果，不进行绝对排除，而是保留语义排序作为回退。最终排序同时考虑检测逻辑与知识的语义关联，以及知识与样例程序结构的适配程度。")
    add_body(doc, "这一过程体现了测试用例在检索阶段的直接作用：测试用例不仅产生自然语言查询所需的检测逻辑，还提供实际程序结构作为第二种检索证据。语义证据回答“该知识能否完成目标操作”，结构证据回答“该知识能否作用于当前代码”。")
    add_body(doc, "仍以条件表达式中的赋值运算为例，语义检索可能同时召回面向普通二元运算、C++重载运算和其他表达式节点的接口。当测试用例AST明确包含BinaryOperator和IfStmt时，与这两个节点直接相关的知识会被优先保留；只适用于当前样例中不存在的节点类型的候选则降低排序。对于hasAncestor、hasParent等通用结构遍历操作，即使其名称不对应单一节点类型，也会因为能够表达条件上下文而继续保留。")
    add_body(doc, "结构重排不是对语义召回的替代。测试用例只能反映有限的代码形式，如果完全按照当前样例节点进行硬过滤，可能删除后续变体所需的知识。因此，方法将AST信息作为相关性增强信号：在有明确结构证据时调整候选顺序，在结构证据不足时仍以语义结果为基础。")

    add_heading(doc, "3.3.4 检索上下文组织与反馈补充", 3)
    add_body(doc, "经过召回和重排后，各检测逻辑获得少量相关知识。检索结果按照模式定位和检查验证两个部分组织：前者包含匹配接口说明及其组合片段，后者包含节点操作方法及验证片段。接口签名与实现片段同时进入生成上下文，使候选生成既有准确的框架接口依据，也有与检测动作对应的实现参考。")
    add_body(doc, "上下文组织需要在完整性和相关性之间取得平衡。如果只返回接口名称，生成模型仍需推断接口之间的关系；如果返回过多实现片段，又容易引入与当前规则无关的条件。方法围绕每条检测逻辑分别选择少量API数据库条目和Meta OP数据库条目，并按照逻辑执行顺序排列，使候选生成能够从目标节点定位逐步过渡到节点验证和诊断输出。")
    add_body(doc, "检索并非只在初始生成时执行。编译反馈可能表明候选代码缺少某项接口知识，测试反馈也可能产生新的检测逻辑。此时，新的知识需求会重新进入API数据库或Meta OP数据库，补充接口签名或实现片段。因而，知识检索与生成、验证之间形成按需补充关系，而不是一次检索后固定不变。")
    add_body(doc, "例如，候选代码在编译时出现绑定节点类型不一致，错误分析可以形成“如何从匹配结果中获取指定类型节点”的新查询。API数据库用于返回节点获取接口及类型信息，Meta OP数据库用于返回“获取节点并在无效时结束检查”的通用片段。若测试阶段出现漏报，则新增的代码模式会先被转化为检测逻辑，再按照相同的分层检索过程获取知识。")
    add_figure(doc, FIGURE_RETRIEVAL, "图3  检测逻辑驱动的分层知识检索方法")

    add_heading(doc, "3.4 阶段3：候选检查器生成与反馈验证", 2)
    add_heading(doc, "3.4.1 多源上下文约束下的候选生成", 3)
    add_body(doc, "候选检查器生成综合使用规则描述、测试用例、样例AST、检测逻辑、检查器结构约束以及检索得到的知识上下文。规则描述保证生成结果与检查目标一致，测试用例和AST提供具体代码证据，检测逻辑规定检查过程，知识上下文限定可使用的框架接口和实现方式。多种信息共同约束候选生成，避免模型仅根据规则名称或单一代码示例推断完整实现。")
    add_body(doc, "生成过程以目标静态分析框架的基本检查器结构为基础，形成包含候选节点匹配和命中后验证的完整检查器。模式定位部分实现目标节点筛选和绑定，检查验证部分实现属性判断、上下文约束、排除条件和诊断输出。Meta OP数据库返回的实现片段作为局部参考，最终仍需根据当前规则统一节点类型、绑定关系和数据传递。")
    add_body(doc, "候选生成按照检测逻辑的先后关系组织内容。首先将目标节点和结构条件转换为候选匹配表达式，并为后续判断需要使用的节点建立一致的绑定；随后在检查阶段取得绑定节点，完成必要的有效性检查和属性读取；最后根据规则约束决定是否输出诊断。API数据库条目用于核对节点类型和接口签名，Meta OP数据库条目用于参考接口组合和控制流程。")
    add_body(doc, "生成结果以完整候选检查器为单位，而不是把检索片段直接拼接后输出。原因在于不同知识条目可能采用不同的局部命名、绑定方式和上下文假设，只有在统一规则语义和结构约束下重新组织，才能保证模式定位与检查验证之间的数据一致性。")

    add_heading(doc, "3.4.2 编译反馈驱动的候选修正", 3)
    add_body(doc, "候选检查器首先接受编译验证，用于检查代码是否满足目标静态分析框架的语法、类型和接口约束。编译失败通常涉及方法签名不匹配、节点类型不一致、依赖缺失或绑定对象使用错误。方法从构建输出中提取有效诊断，结合当前候选代码分析错误原因，并形成修正方向。")
    add_body(doc, "当错误分析表明现有上下文缺少相关接口或实现方式时，系统根据新的知识需求进行补充检索，再利用错误诊断和补充知识生成修正版本。编译反馈主要解决候选检查器的可执行性问题，使后续测试验证集中于规则行为，而不受基础接口错误干扰。")
    add_body(doc, "编译修正保持原有检测目标不变，重点调整接口选择、节点类型、依赖关系和调用形式。如果错误来自单个API使用不当，优先从API数据库补充准确签名；如果错误涉及绑定、节点获取或多个接口之间的衔接，则同时从Meta OP数据库获取相应组合方式。修正后的候选重新接受编译验证，直到满足目标静态分析框架的基本约束或达到本轮停止条件。")

    add_heading(doc, "3.4.3 测试反馈驱动的逻辑增强", 3)
    add_body(doc, "编译通过后，候选检查器在测试套件上运行。违规用例未产生预期告警时判定为漏报，说明检测逻辑没有覆盖该代码形式，或者当前约束过严；合规用例产生告警时判定为误报，说明匹配范围过宽或缺少排除条件。两类失败对应不同的逻辑修正方向。")
    add_body(doc, "对于漏报，方法从失败用例中提取新的代码模式，补充目标节点、结构关系或上下文路径；对于误报，方法比较失败用例与已通过用例的差异，增加类型、作用域、上下文或语义排除条件。更新后的检测逻辑再次触发知识检索和候选生成，使测试失败能够转化为下一轮生成所需的明确约束。")
    add_body(doc, "每次局部修正后均重新运行测试套件，而不是只验证当前失败用例。全量回归用于判断新增匹配条件是否引入新的误报，以及新增排除条件是否造成新的漏报。已经通过的测试用例因此构成后续修改需要保持的行为边界。")
    add_body(doc, "测试反馈与编译反馈的作用层次不同。编译反馈主要判断实现是否符合框架约束，测试反馈主要判断检测语义是否符合规则要求。前者通常直接修正候选代码，后者通常先调整检测逻辑和检查器结构约束，再重新执行知识检索和候选生成。两类反馈分开处理，可以避免仅通过修改代码表面形式掩盖规则逻辑中的缺失。")

    add_heading(doc, "3.4.4 迭代过程与结果输出", 3)
    add_body(doc, "候选生成、编译验证、测试验证和逻辑增强共同构成迭代过程。当候选检查器通过编译，且测试结果满足违规用例告警、合规用例静默的要求时，当前版本作为最终检查器输出。若某一候选在限定轮次内不能满足要求，则更换种子用例或记录尚未解决的测试情形，避免生成过程无界执行。")
    add_body(doc, "从方法整体看，三个阶段分别解决了规则如何转化为程序结构约束、检测逻辑如何获得目标静态分析框架知识、候选代码如何通过确定性反馈逐步满足测试要求三个问题。知识库提供框架层面的先验，测试用例提供规则层面的证据，编译器和测试套件共同构成候选结果的外部判据。")


def build():
    if not DOCUMENT.exists():
        raise FileNotFoundError(DOCUMENT)
    make_kb_figure()
    make_retrieval_figure()

    doc = Document(DOCUMENT)
    chapter_index = next(i for i, p in enumerate(doc.paragraphs)
                         if p.text.strip().startswith("3 测试用例驱动"))
    chapter_heading = doc.paragraphs[chapter_index]
    # Remove all old Chapter 3 body elements, including tables and drawings.
    body = doc._element.body
    element = chapter_heading._p.getnext()
    while element is not None and element.tag != qn("w:sectPr"):
        next_element = element.getnext()
        body.remove(element)
        element = next_element

    chapter_heading.text = "3 测试用例驱动的静态代码检查器生成技术"
    chapter_heading.style = "Normal"
    format_heading(chapter_heading, 1)
    write_chapter(doc)
    add_page_numbers(doc)

    # Write atomically while keeping the user-requested filename unchanged.
    fd, temp_name = tempfile.mkstemp(prefix=".autochecker-ch3-", suffix=".docx",
                                     dir=str(DOCUMENT.parent))
    os.close(fd)
    try:
        doc.save(temp_name)
        os.replace(temp_name, DOCUMENT)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    print(DOCUMENT)


if __name__ == "__main__":
    build()
